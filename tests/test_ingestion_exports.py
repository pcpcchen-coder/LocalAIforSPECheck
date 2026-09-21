import copy
import hashlib
import io
import json
import re
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter

from spec_check.exports import export_run
from spec_check.ingestion import MAX_BLOCK_CHARS, extract_document


def test_text_bom_locations_and_long_line_are_not_truncated(tmp_path):
    path = tmp_path / "spec.txt"
    original = "  第一項：電壓 220V  " + "規格" * 1800
    path.write_text("標題\n\n" + original + "\n尾項", encoding="utf-8-sig")
    result = extract_document(path, path.name, "standard")
    assert result["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert result["blocks"][0]["text"] == "標題"
    source_line = [block for block in result["blocks"] if "第 3 行" in block["location"]]
    assert "".join(block["text"] for block in source_line) == original
    assert all(len(block["text"]) <= MAX_BLOCK_CHARS for block in result["blocks"])
    assert [block["id"] for block in result["blocks"]] == [f"B{i:05d}" for i in range(1, len(result["blocks"]) + 1)]
    assert result["blocks"][-1]["text"] == "尾項"


def test_csv_multiline_and_bom(tmp_path):
    path = tmp_path / "spec.csv"
    path.write_text('名稱,要求\n"電源,主機","第一行\n第二行"\n', encoding="utf-8-sig")
    result = extract_document(path, "規格.CSV", "product")
    assert len(result["blocks"]) == 2
    assert "電源,主機" in result["blocks"][1]["text"]
    assert "第一行\n第二行" in result["blocks"][1]["text"]
    assert "第 2–3 行" in result["blocks"][1]["location"]


def test_docx_keeps_paragraphs_tables_and_header(tmp_path):
    path = tmp_path / "spec.docx"
    document = Document()
    document.add_heading("產品規格", 1)
    document.add_paragraph("環境溫度 -10 至 60°C")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "項目"
    table.cell(0, 1).text = "規格"
    table.cell(1, 0).text = "電壓"
    table.cell(1, 1).text = "220 V"
    document.sections[0].header.paragraphs[0].text = "機密設計版本 A"
    document.sections[0].footer.paragraphs[0].text = "2026 年版"
    document.save(path)
    result = extract_document(path, path.name, "product")
    texts = "\n".join(block["text"] for block in result["blocks"])
    for expected in ["產品規格", "環境溫度", "220 V", "機密設計版本 A", "2026 年版"]:
        assert expected in texts
    assert any("表格 1 第 2 列" in block["location"] for block in result["blocks"])
    table_row = next(block for block in result["blocks"] if "表格 1 第 2 列" in block["location"])
    assert "推定表頭原文" in table_row["text"] and "項目" in table_row["text"] and "規格" in table_row["text"]
    assert result["warnings"]


def test_xlsx_formula_without_cache_is_visible_and_warned(tmp_path):
    path = tmp_path / "spec.xlsx"
    workbook = Workbook()
    workbook.active.title = "額定"
    workbook.active.append(["項目", "数值"])
    workbook.active.append(["電壓", 220])
    workbook.active["B3"] = "=B2*2"
    workbook.active["B4"] = 0.95
    workbook.active["B4"].number_format = "0.0%"
    hidden = workbook.create_sheet("備註")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "隱藏規格"
    workbook.save(path)
    result = extract_document(path, path.name, "standard")
    texts = "\n".join(block["text"] for block in result["blocks"])
    assert "=B2*2" in texts
    assert "公式結果未快取" in texts
    assert "隱藏規格" in texts
    assert "0.95 [顯示格式：0.0%]" in texts
    assert any("B3" in warning and "快取" in warning for warning in result["warnings"])
    assert any("隱藏工作表" in warning for warning in result["warnings"])
    data_row = next(block for block in result["blocks"] if "「額定」第 2 列" in block["location"])
    assert "推定表頭原文" in data_row["text"] and "項目" in data_row["text"]


def _rewrite_sheet_xml(path, transform):
    with zipfile.ZipFile(path) as archive:
        items = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with zipfile.ZipFile(path, "w") as archive:
        for item, data in items:
            if item.filename == "xl/worksheets/sheet1.xml":
                data = transform(data.decode("utf-8")).encode("utf-8")
            archive.writestr(item, data)


@pytest.mark.parametrize("dimension", ["A1:A1", "A1:XFD1048576", "B3:B3"])
def test_xlsx_actual_rows_override_incorrect_dimension(tmp_path, dimension):
    path = tmp_path / "incorrect-dimension.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "項目"
    workbook.active["A2"] = "不得漏掉第二列"
    workbook.active["B3"] = "=220+10"
    workbook.save(path)

    def transform(xml):
        xml = re.sub(r'<dimension ref="[^"]+"\s*/>', f'<dimension ref="{dimension}"/>', xml)
        return xml.replace("<f>220+10</f><v></v>", "<f>220+10</f><v>230</v>")

    _rewrite_sheet_xml(path, transform)
    result = extract_document(path, path.name, "standard")
    text = "\n".join(block["text"] for block in result["blocks"])
    assert "不得漏掉第二列" in text
    assert "B3：230 [公式：=220+10]" in text
    assert any("第 3 列" in block["location"] for block in result["blocks"])
    assert any("宣告範圍" in warning for warning in result["warnings"])
    assert not any("沒有可讀取的快取" in warning for warning in result["warnings"])


def test_xlsx_actual_range_limit_applies_despite_small_dimension(tmp_path):
    path = tmp_path / "actual-too-large.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "項目"
    workbook.active["B500001"] = "超出範圍"
    workbook.save(path)
    _rewrite_sheet_xml(path, lambda xml: re.sub(r'<dimension ref="[^"]+"\s*/>', '<dimension ref="A1:A1"/>', xml))
    with pytest.raises(ValueError, match="實際使用範圍超過 100 萬格"):
        extract_document(path, path.name, "standard")


def test_blank_pdf_does_not_claim_extraction_complete(tmp_path):
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with path.open("wb") as stream:
        writer.write(stream)
    result = extract_document(path, path.name, "standard")
    assert result["blocks"] == []
    assert any("第 1 頁沒有可抽取文字" in warning for warning in result["warnings"])
    assert any("未 OCR" in warning for warning in result["warnings"])


def test_text_pdf_preserves_page_location(tmp_path):
    # Construct a text-layer PDF without adding a reportlab runtime dependency.
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    path = tmp_path / "text.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 800 Td (Voltage 220 V) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as target:
        writer.write(target)
    result = extract_document(path, path.name, "product")
    assert "Voltage 220 V" in result["blocks"][0]["text"]
    assert result["blocks"][0]["location"] == "PDF 第 1 頁 第 1–1 行"


@pytest.mark.parametrize("filename,content,error", [("bad.xls", b"x", "不支援"), ("bad.docx", b"not zip", "有效"), ("bad.txt", b"\xff", "UTF-8")])
def test_unsupported_or_invalid_documents_are_rejected(tmp_path, filename, content, error):
    path = tmp_path / filename
    path.write_bytes(content)
    with pytest.raises(ValueError, match=error):
        extract_document(path, filename, "product")


def test_archive_bomb_guard(tmp_path):
    path = tmp_path / "huge.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "0" * 2_000_000)
    with pytest.raises(ValueError, match="壓縮比例"):
        extract_document(path, path.name, "product")


@pytest.fixture
def run_snapshot():
    return {
        "id": "run-test", "project_id": "p1", "status": "completed", "mode": "local",
        "created_at": "2026-09-21T00:00:00Z", "finished_at": "2026-09-21T00:01:00Z", "total": 2, "completed": 2,
        "settings": {"model": "local-model", "api_key": "TOP-SECRET", "base_url": "http://127.0.0.1:1234/v1"},
        "documents": [
            {"id": "product-1", "name": "產品.txt", "role": "product", "sha256": "abc", "extraction_confirmed": True, "warnings": ["需人工確認"], "blocks": [{"id": "B00001", "location": "文字第 1 行", "text": "額定電壓 220 V"}, {"id": "B00002", "location": "文字第 2 行", "text": "所有未引述原文也須保留"}]},
            {"id": "standard-1", "name": "規範.txt", "role": "standard", "sha256": "def", "extraction_confirmed": True, "warnings": [], "blocks": [{"id": "B00001", "location": "文字第 1 行", "text": "額定電壓 230 V"}, {"id": "B00002", "location": "文字第 2 行", "text": "防水 IP67"}]},
        ],
        "results": [
            {"id": "r1", "run_id": "run-test", "standard_id": "standard-1", "standard_name": "規範.txt", "block_id": "B00001", "location": "文字第 1 行", "requirement": '<script>alert("x")</script>', "status": "mismatch", "explanation": "額定電壓不同", "differences": ["差 10 V"], "confidence": 0.9, "evidence": [{"block_id": "B00001", "location": "文字第 1 行", "quote": "額定電壓 220 V"}], "product_coverage": {"scanned": 2, "total": 2}, "warnings": [], "review": {"decision": "changed", "final_status": "partial", "reviewer": "工程師", "note": "依實測確認", "version": 1, "updated_at": "2026-09-21T01:00:00Z"}, "history": [{"id": "ev1", "result_id": "r1", "decision": "changed", "final_status": "partial", "reviewer": "工程師", "note": "依實測確認", "created_at": "2026-09-21T01:00:00Z", "version": 1}]},
            {"id": "r2", "standard_id": "standard-1", "standard_name": "規範.txt", "block_id": "B00002", "location": "文字第 2 行", "requirement": "防水 IP67", "status": "missing", "explanation": "產品文件未載明", "differences": ["待取得測試報告"], "evidence": [], "confidence": 0.8, "product_coverage": {"scanned": 2, "total": 2}, "warnings": [], "review": {"decision": "pending", "version": 0}, "history": []},
        ],
        "rankings": [{"standard_id": "standard-1", "standard_name": "規範.txt", "score": 25.0, "total": 2, "completed": 2, "reviewed": 1, "counts": {"partial": 1, "missing": 1}, "coverage": 1.0}],
    }


def test_json_preserves_complete_snapshot_without_mutating_or_secrets(run_snapshot):
    original = copy.deepcopy(run_snapshot)
    run_snapshot["extra"] = {"Authorization": "ANOTHER-SECRET", "nested": [{"api-key": "HIDDEN"}]}
    data, mime, filename = export_run(run_snapshot, "json")
    report = json.loads(data)
    assert "TOP-SECRET" not in data.decode()
    assert "ANOTHER-SECRET" not in data.decode()
    assert "HIDDEN" not in data.decode()
    assert report["results"] == original["results"]
    assert report["documents"] == original["documents"]
    assert run_snapshot["settings"]["api_key"] == "TOP-SECRET"
    assert mime.startswith("application/json")
    assert filename == "spec-check-run-test.json"


def test_html_is_offline_escaped_and_keeps_pending_source_and_history(run_snapshot):
    data, mime, _ = export_run(run_snapshot, "html")
    text = data.decode()
    assert "<script>" not in text
    assert "&lt;script&gt;" in text
    assert 'src="http' not in text
    assert "所有未引述原文也須保留" in text
    assert "ev1" in text and "依實測確認" in text
    assert "防水 IP67" in text and "待覆核" in text
    assert "TOP-SECRET" not in text
    assert "產品文件未載明不等於產品本身不合規" in text
    assert mime.startswith("text/html")


def test_xlsx_formulas_are_plain_text_and_long_content_is_recoverable(run_snapshot):
    injection = '=HYPERLINK("https://invalid.example", "open")'
    run_snapshot["results"][0]["requirement"] = injection
    run_snapshot["results"][0]["explanation"] = "Z" * 70_000
    run_snapshot["results"][0]["review"]["note"] = "control\x01character"
    data, mime, _ = export_run(run_snapshot, "xlsx")
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    assert workbook.sheetnames == ["Results", "History", "Documents", "Run"]
    results = workbook["Results"]
    assert results["F2"].value == injection
    assert results["F2"].data_type == "s"
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    assert "".join(results.cell(row, 10).value or "" for row in range(2, 5)) == "Z" * 70_000
    complete_json = "".join(results.cell(row, 21).value or "" for row in range(2, 5))
    assert json.loads(complete_json)["explanation"] == "Z" * 70_000
    assert results["R2"].value == "control\\u0001character"
    all_cells = "\n".join(str(cell.value or "") for sheet in workbook for row in sheet for cell in row)
    assert "TOP-SECRET" not in all_cells
    assert "所有未引述原文也須保留" in all_cells
    assert "ev1" in all_cells
    assert mime.endswith("spreadsheetml.sheet")


def test_exports_reject_unknown_format_and_sanitize_filename(run_snapshot):
    with pytest.raises(ValueError, match="匯出格式"):
        export_run(run_snapshot, "exe")
    run_snapshot["id"] = "../../\r\nX-Evil:yes"
    _, _, name = export_run(run_snapshot, "json")
    assert "/" not in name and "\r" not in name and "\n" not in name
