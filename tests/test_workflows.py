"""v0.3 API lifecycle acceptance tests, using synthetic local model work.

These exercise the real HTTP API, SQLite persistence, jobs, immutable document
snapshots, reviews and export archives. The model functions are deterministic
fixtures; passing this suite is not a real-document accuracy benchmark.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import threading
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from spec_check.app import create_app

BASE_URL = "http://127.0.0.1:8765"
PREFIX = "/api/v2"
REVIEWER = "合成驗收工程師"


def ok(response, status=200):
    assert response.status_code == status, response.text
    return response.json()


def wait(client, path, statuses=("completed",), timeout=15):
    deadline = time.monotonic() + timeout
    state = None
    while time.monotonic() < deadline:
        state = ok(client.get(PREFIX + path))
        if state["status"] in statuses:
            return state
        if state["status"] == "failed" and "failed" not in statuses:
            pytest.fail(f"Background job failed: {state}")
        time.sleep(0.01)
    pytest.fail(f"Did not reach {statuses}: {state}")


def upload(client, name, text, role="standard"):
    route = "/products" if role == "product" else "/library"
    return ok(client.post(PREFIX + route, files={"file": (name, text.encode(), "text/plain")}))


def doc(client, identifier):
    return ok(client.get(f"{PREFIX}/documents/{identifier}"))


def rows(client, path, **params):
    return ok(client.get(PREFIX + path, params=params))["items"]


def item_edit(item, **changes):
    editable = {key: value for key, value in item.items() if key in {"name", "parameter", "value", "unit", "operator", "conditions", "exceptions", "test_method", "criticality", "criticality_basis", "quote", "block_id", "kind"}}
    return dict(editable, **changes)


def mutation(document, **kwargs):
    return dict(expected_version=document["version"], reviewer=REVIEWER, note="合成驗收：已核對原文。", **kwargs)


def extract(client, document):
    job = ok(client.post(f"{PREFIX}/documents/{document['id']}/extract", json={}))
    wait(client, f"/jobs/{job['id']}")
    return doc(client, document["id"])


def confirm(client, document):
    return ok(client.post(f"{PREFIX}/documents/{document['id']}/confirm", json=mutation(document, acknowledge_warnings=True)))


def prepare(client, name, text, role="standard"):
    return confirm(client, extract(client, upload(client, name, text, role)))


def analyse(client, product, standards, name="合成分析"):
    created = ok(client.post(PREFIX + "/analyses", json={
        "name": name, "product_id": product["id"], "library_ids": [d["id"] for d in standards],
        "context": {"purpose": "合成儲能控制器", "environment": "室內", "market": "未知", "notes": "僅軟體驗收"},
        "comparison_mode": "focused",
    }))
    return wait(client, f"/analyses/{created['id']}", ("awaiting_selection",))


def compare(client, analysis):
    started = ok(client.post(f"{PREFIX}/analyses/{analysis['id']}/compare", json={"expected_version": analysis["version"]}))
    return wait(client, f"/analyses/{started['id']}")


@pytest.fixture
def model(monkeypatch):
    # Import after collection so this file can be prepared while the service is
    # being implemented. Only model boundary work is replaced in lifecycle tests.
    from spec_check import analysis_engine as ae

    state = {"extract": [], "screen": [], "compare": []}

    def extracted(block, role, settings, **kwargs):
        state["extract"].append(copy.deepcopy(block))
        text = block["text"].strip()
        item = {
            "name": text.split("：", 1)[0], "parameter": text.split("：", 1)[0],
            "value": text.split("：", 1)[-1], "unit": "", "operator": "=",
            "conditions": "", "exceptions": "", "test_method": "",
            "criticality": "high" if "安全" in text else "unknown", "criticality_basis": "",
            "quote": text, "block_id": block["id"],
            "kind": "specification" if role == "product" else "requirement",
        }
        return {"items": [item], "warnings": [], "coverage": "complete"}

    def screened(product, standard, context, settings, **kwargs):
        state["screen"].append(standard["id"])
        return {"relevance": "low" if "無關" in standard["name"] else "high", "relevance_score": 10 if "無關" in standard["name"] else 90,
                "applicability": "unknown", "reason": "合成初篩：適用條件需人工確認。", "evidence": [], "warnings": []}

    def compared(item, product, standard_context, settings, **kwargs):
        state["compare"].append(copy.deepcopy(item))
        product_items = product.get("items", []) if isinstance(product, dict) else product
        product_blocks = product.get("blocks", []) if isinstance(product, dict) else []
        first = product_items[0] if product_items else {}
        block = product_blocks[0] if product_blocks else first
        result = {"status": "partial", "explanation": "合成判定：需補測試證據。", "differences": ["缺少測試報告。"],
                "evidence": [{"block_id": block.get("id", first.get("block_id", "B00001")), "location": block.get("location", "第 1 行"), "quote": block.get("text", first.get("quote", "額定電壓：48 V"))}],
                "confidence": 0.7, "warnings": [], "product_coverage": {"scanned": 1, "total": 1},
                "matched_product_item_ids": [first["id"]] if first.get("id") else [],
                "retrieval": {"mode": "focused", "candidate_count": len(product_items), "total_product_items": len(product_items), "total": len(product_items), "complete": True}}
        result["risk"] = ae.assess_risk(result, item)
        return result

    monkeypatch.setattr(ae, "extract_block", extracted)
    monkeypatch.setattr(ae, "screen_standard", screened)
    monkeypatch.setattr(ae, "compare_item", compared)
    return state


@pytest.fixture
def client(tmp_path, model):
    with TestClient(create_app(tmp_path), base_url=BASE_URL) as current:
        ok(current.put("/api/settings", json={"model": "synthetic-lifecycle-model"}))
        yield current


def test_source_confirmation_item_edit_split_versions_and_audit(client):
    product = upload(client, "產品.txt", "額定電壓：48 V\n", "product")
    standard = upload(client, "規範.txt", "額定電壓：48 V；工作溫度：0–40 C\n")
    assert client.post(PREFIX + "/analyses", json={"name": "未確認", "product_id": product["id"], "library_ids": [standard["id"]]}).status_code == 400
    standard = extract(client, standard)
    extracted = rows(client, f"/documents/{standard['id']}/items")
    assert len(extracted) == 1
    item = extracted[0]
    url = f"{PREFIX}/documents/{standard['id']}/items/{item['id']}"
    invalid = item_edit(item, quote="不存在的 900 V 來源")
    assert client.put(url, json=mutation(standard, item=invalid)).status_code == 400
    assert doc(client, standard["id"])["version"] == standard["version"]
    updated = item_edit(item, name="合成複合規格")
    ok(client.put(url, json=mutation(standard, item=updated)))
    assert client.put(url, json=mutation(standard, item=updated)).status_code == 409
    standard = doc(client, standard["id"])
    fragments = [item_edit(item, name="額定電壓", parameter="額定電壓", quote="額定電壓：48 V"),
                 item_edit(item, name="工作溫度", parameter="工作溫度", quote="工作溫度：0–40 C")]
    ok(client.post(url + "/split", json=mutation(standard, items=fragments)))
    standard = doc(client, standard["id"])
    items = rows(client, f"/documents/{standard['id']}/items")
    assert len(items) >= 2
    assert {"額定電壓", "工作溫度"}.issubset({i["name"] for i in items})
    assert len({i["id"] for i in items}) == len(items)
    assert all(i["quote"] in standard["blocks"][0]["text"] for i in items)
    standard = confirm(client, standard)
    assert standard["confirmed"]
    events = rows(client, f"/documents/{standard['id']}/audit")
    assert len(events) >= 4
    assert REVIEWER in json.dumps(events, ensure_ascii=False)


def test_three_hundred_standards_paged_screened_without_silent_exclusion(client, model):
    product = prepare(client, "產品.txt", "額定電壓：48 V\n", "product")
    standards = [upload(client, f"合成標準-{n:03d}.txt", f"額定電壓：{n + 24} V\n") for n in range(300)]
    job = ok(client.post(PREFIX + "/library/extract", json={"document_ids": [d["id"] for d in standards]}))
    wait(client, f"/jobs/{job['id']}", timeout=45)
    summaries = []
    for offset in (0, 100, 200):
        page = ok(client.get(PREFIX + "/library", params={"offset": offset, "limit": 100}))
        assert page["total"] == 300
        assert len(page["items"]) == 100
        assert all(not {"blocks", "items"}.intersection(d) for d in page["items"])
        summaries.extend(page["items"])
    assert len({d["id"] for d in summaries}) == 300
    ok(client.post(PREFIX + "/library/confirm", json={"documents": [{"id": d["id"], "expected_version": d["version"]} for d in summaries], "reviewer": REVIEWER, "note": "已檢查本測試產生的 300 份合成規範。", "acknowledge_warnings": True}))
    analysis = analyse(client, product, [])  # Empty selection means all library documents.
    scanned = []
    for offset in (0, 100, 200):
        page = ok(client.get(f"{PREFIX}/analyses/{analysis['id']}/screening", params={"offset": offset, "limit": 100}))
        assert page["total"] == 300
        scanned.extend(page["items"])
    assert {d["standard_id"] for d in scanned} == {d["id"] for d in standards}
    assert set(model["screen"]) == {d["id"] for d in standards}
    assert all(s["selected"] for s in scanned)
    assert len(model["screen"]) == 300
    assert not {"documents", "results", "screenings", "settings"}.intersection(analysis)
    assert len(json.dumps(analysis)) < 30000


def test_selection_requires_reason_cas_and_preserves_every_screening_row(client):
    product = prepare(client, "產品.txt", "額定電壓：48 V\n", "product")
    standards = [prepare(client, name, text) for name, text in [("相關.txt", "額定電壓：48 V\n"), ("無關.txt", "航空高度：8000 m\n")]]
    analysis = analyse(client, product, standards)
    initial = rows(client, f"/analyses/{analysis['id']}/screening")
    assert len(initial) == 2 and all(s["selected"] for s in initial)
    low = next(s for s in initial if s["standard_id"] == standards[1]["id"])
    assert low["selected"], "Low model relevance must not silently exclude a standard"
    url = f"{PREFIX}/analyses/{analysis['id']}/selection"
    decision = {"standard_id": standards[1]["id"], "selected": False, "reason": ""}
    assert client.post(url, json=dict(mutation(analysis, decisions=[decision]), note="")).status_code == 400
    decision["reason"] = "合成驗收：此產品不屬於航空設備。"
    ok(client.post(url, json=mutation(analysis, decisions=[decision])))
    assert client.post(url, json=mutation(analysis, decisions=[decision])).status_code == 409
    current = ok(client.get(f"{PREFIX}/analyses/{analysis['id']}"))
    finished = compare(client, current)
    screening = rows(client, f"/analyses/{analysis['id']}/screening")
    assert len(screening) == 2
    result = rows(client, f"/analyses/{analysis['id']}/results")
    assert {r["standard_id"] for r in result} == {standards[0]["id"]}
    audit = rows(client, f"/analyses/{analysis['id']}/audit")
    assert decision["reason"] in json.dumps(audit, ensure_ascii=False)
    assert finished["status"] == "completed"
    assert finished["coverage"]["selected_scope_complete"]
    assert not finished["coverage"]["coverage_complete"], "Excluded standards prevent all-library completeness"
    extras = ok(client.get(f"{PREFIX}/analyses/{analysis['id']}/product-extras"))
    assert not extras["coverage_complete"]


def test_analysis_snapshot_does_not_follow_later_source_edits(client):
    product = prepare(client, "產品.txt", "額定電壓：48 V\n", "product")
    standard = prepare(client, "標準.txt", "額定電壓：48 V\n")
    analysis = analyse(client, product, [standard])
    before = ok(client.get(f"{PREFIX}/analyses/{analysis['id']}/export?format=json"))
    source = doc(client, standard["id"])
    item = rows(client, f"/documents/{source['id']}/items")[0]
    ok(client.put(f"{PREFIX}/documents/{source['id']}/items/{item['id']}", json=mutation(source, item=item_edit(item, name="之後人工修正的名稱"))))
    after = ok(client.get(f"{PREFIX}/analyses/{analysis['id']}/export?format=json"))
    assert before["documents"] == after["documents"]
    assert "之後人工修正的名稱" not in json.dumps(after["documents"], ensure_ascii=False)
    assert not doc(client, source["id"])["confirmed"], "Editing a confirmed source requires confirmation again"


def test_risk_review_conflicts_history_and_immutable_export_archives(client):
    product = prepare(client, "產品.txt", "額定電壓：48 V\n", "product")
    standard = prepare(client, "安全規範.txt", "安全耐壓：500 V\n")
    analysis = compare(client, analyse(client, product, [standard]))
    result = rows(client, f"/analyses/{analysis['id']}/results")[0]
    detail_url = f"{PREFIX}/analyses/{analysis['id']}/results/{result['id']}"
    body = {"expected_version": result["review"]["version"], "reviewer": REVIEWER,
            "decision": "changed", "final_status": "uncertain", "risk_level": "high", "note": ""}
    assert client.post(detail_url + "/reviews", json=body).status_code == 400
    body["note"] = "合成驗收：安全測試證據不足，需優先補測。"
    changed = ok(client.post(detail_url + "/reviews", json=body))
    assert changed["status"] == result["status"], "Human review cannot rewrite the AI verdict"
    assert changed["review"]["risk_level"] == "high"
    assert changed["review"]["final_status"] == "uncertain"
    assert client.post(detail_url + "/reviews", json=body).status_code == 409
    assert len(ok(client.get(detail_url))["history"]) == 1
    snapshots = {}
    for fmt in ("json", "html", "xlsx"):
        response = client.get(f"{PREFIX}/analyses/{analysis['id']}/export", params={"format": fmt})
        assert response.status_code == 200, response.text
        snapshots[fmt] = response.content
        if fmt == "json":
            exported = response.json()
            assert exported["schema_version"] == 2 and exported["kind"] == "analysis"
            assert exported["results"][0]["review"]["risk_level"] == "high"
            assert {"documents", "screenings", "product_extras", "audit", "coverage"}.issubset(exported)
        elif fmt == "html":
            assert body["note"] in response.text
        else:
            book = load_workbook(io.BytesIO(response.content), read_only=True)
            assert len(book.sheetnames) >= 6
            assert any(body["note"] in str(value) for sheet in book for row in sheet.iter_rows(values_only=True) for value in row)
            book.close()
    archives = rows(client, f"/analyses/{analysis['id']}/exports")
    assert len(archives) == 3
    reopened = ok(client.post(detail_url + "/reviews", json={
        "expected_version": changed["review"]["version"], "decision": "reopened", "reviewer": REVIEWER,
        "note": "收到新測試資料，重新確認。"}))
    assert reopened["review"]["final_status"] is None
    assert len(reopened["history"]) == 2
    for archive in archives:
        response = client.get(f"/api/exports/{archive['id']}/download")
        assert response.status_code == 200
        assert response.content == snapshots[archive["format"]]
        assert hashlib.sha256(response.content).hexdigest() == archive["sha256"]


def test_cancel_waiting_job_and_resume_does_not_duplicate_results(client, model, monkeypatch):
    from spec_check import analysis_engine as ae
    from spec_check.engine import ComparisonCancelled
    product = prepare(client, "產品.txt", "額定電壓：48 V\n", "product")
    standard = prepare(client, "標準.txt", "額定電壓：48 V\n")
    analysis = analyse(client, product, [standard])
    entered, release = threading.Event(), threading.Event()
    original = ae.compare_item

    def blocked(*args, **kwargs):
        entered.set()
        assert release.wait(8)
        cancel = kwargs.get("cancel_check")
        if cancel and cancel():
            raise ComparisonCancelled("合成取消")
        return original(*args, **kwargs)

    monkeypatch.setattr(ae, "compare_item", blocked)
    ok(client.post(f"{PREFIX}/analyses/{analysis['id']}/compare", json={"expected_version": analysis["version"]}))
    try:
        assert entered.wait(5)
        progress = ok(client.get(f"{PREFIX}/analyses/{analysis['id']}"))
        assert not {"documents", "results", "settings"}.intersection(progress)
        ok(client.post(f"{PREFIX}/analyses/{analysis['id']}/cancel"))
    finally:
        release.set()
    wait(client, f"/analyses/{analysis['id']}", ("cancelled",))
    assert rows(client, f"/analyses/{analysis['id']}/results") == []
    monkeypatch.setattr(ae, "compare_item", original)
    ok(client.post(f"{PREFIX}/analyses/{analysis['id']}/resume"))
    wait(client, f"/analyses/{analysis['id']}")
    result = rows(client, f"/analyses/{analysis['id']}/results")
    assert len(result) == 1


def test_restart_marks_interrupted_and_resumes_saved_checkpoint(tmp_path, model):
    # Simulate abrupt process loss by retaining complete result checkpoints but
    # leaving the persisted worker status active. No live worker races this edit.
    with TestClient(create_app(tmp_path), base_url=BASE_URL) as first:
        ok(first.put("/api/settings", json={"model": "synthetic-lifecycle-model"}))
        product = prepare(first, "產品.txt", "額定電壓：48 V\n", "product")
        standard = prepare(first, "標準.txt", "額定電壓：48 V\n")
        analysis = compare(first, analyse(first, product, [standard]))
        original = rows(first, f"/analyses/{analysis['id']}/results")
        calls = len(model["compare"])
        store = first.app.state.store
        saved = store.get("v2_analysis", analysis["id"])
        saved["status"] = "comparing"
        job = store.get("v2_job", saved["job_id"])
        job["status"] = "running"
        store.put_many([("v2_analysis", saved, None), ("v2_job", job, saved["id"])])
    with TestClient(create_app(tmp_path), base_url=BASE_URL) as second:
        interrupted = ok(second.get(f"{PREFIX}/analyses/{analysis['id']}"))
        assert interrupted["status"] == "interrupted"
        assert interrupted["progress"]["request_started_at"] is None
        assert rows(second, f"/analyses/{analysis['id']}/results") == original
        ok(second.post(f"{PREFIX}/analyses/{analysis['id']}/resume"))
        wait(second, f"/analyses/{analysis['id']}")
        assert rows(second, f"/analyses/{analysis['id']}/results") == original
        assert len(model["compare"]) == calls, "Already committed items must not be recomputed"


def test_model_failure_retains_uncertain_row_and_never_exposes_api_key(client, monkeypatch):
    from spec_check import analysis_engine as ae
    secret = "synthetic-private-token-must-not-leak"
    ok(client.put("/api/settings", json={"api_key": secret}))
    product = prepare(client, "產品.txt", "額定電壓：48 V\n", "product")
    standard = prepare(client, "標準.txt", "安全耐壓：500 V\n")
    analysis = analyse(client, product, [standard])

    def failure(*args, **kwargs):
        raise RuntimeError("Model leaked request key " + secret)

    monkeypatch.setattr(ae, "compare_item", failure)
    analysis = compare(client, analysis)
    result = rows(client, f"/analyses/{analysis['id']}/results")
    assert len(result) == 1
    assert result[0]["status"] == "uncertain"
    assert not result[0]["evidence"]
    assert not result[0]["retrieval"]["complete"]
    for path in (f"/analyses/{analysis['id']}", f"/analyses/{analysis['id']}/results",
                 f"/analyses/{analysis['id']}/audit", f"/documents/{product['id']}",
                 f"/documents/{standard['id']}/audit"):
        response = client.get(PREFIX + path)
        assert response.status_code == 200
        assert secret not in response.text
    for fmt in ("json", "html"):
        response = client.get(f"{PREFIX}/analyses/{analysis['id']}/export", params={"format": fmt})
        assert response.status_code == 200
        assert secret not in response.text


def test_partial_batch_cancel_resume_preserves_completed_confirmed_document(client, model, monkeypatch):
    from spec_check import analysis_engine as ae
    from spec_check.engine import ComparisonCancelled
    documents = [upload(client, f"批次-{n}.txt", f"額定電壓：{n + 50} V\n") for n in range(3)]
    original = ae.extract_block
    entered, release = threading.Event(), threading.Event()
    blocked_once = False

    def blocked(block, *args, **kwargs):
        nonlocal blocked_once
        if "51 V" in block["text"] and not blocked_once:
            blocked_once = True
            entered.set()
            assert release.wait(8)
            if kwargs.get("cancel_check") and kwargs["cancel_check"]():
                raise ComparisonCancelled("合成取消抽取")
        return original(block, *args, **kwargs)

    monkeypatch.setattr(ae, "extract_block", blocked)
    job = ok(client.post(PREFIX + "/library/extract", json={"document_ids": [d["id"] for d in documents]}))
    try:
        assert entered.wait(5)
        first = confirm(client, doc(client, documents[0]["id"]))
        assert first["confirmed"]
        ok(client.post(f"{PREFIX}/jobs/{job['id']}/cancel"))
    finally:
        release.set()
    wait(client, f"/jobs/{job['id']}", ("cancelled",))
    ok(client.post(f"{PREFIX}/jobs/{job['id']}/resume"))
    finished = wait(client, f"/jobs/{job['id']}")
    current = doc(client, documents[0]["id"])
    assert current["confirmed"] and current["version"] == first["version"]
    assert sum("50 V" in b["text"] for b in model["extract"]) == 1
    assert finished["completed"] == finished["total"] == 3
    assert all(doc(client, d["id"])["index_status"] == "ready" for d in documents)


def test_item_source_cannot_move_to_another_block_with_identical_text(client):
    source = extract(client, upload(client, "重複原文.txt", "額定電壓：48 V\n額定電壓：48 V\n"))
    assert len(source["blocks"]) == 2
    items = rows(client, f"/documents/{source['id']}/items")
    item = items[0]
    other = next(b for b in source["blocks"] if b["id"] != item["block_id"])
    response = client.put(f"{PREFIX}/documents/{source['id']}/items/{item['id']}", json=mutation(source, item=item_edit(item, block_id=other["id"])))
    assert response.status_code == 400
    assert rows(client, f"/documents/{source['id']}/items") == items


def test_import_classic_project_reuses_documents_and_keeps_original_history(client):
    project = ok(client.post("/api/projects", json={"name": "既有使用者專案"}))
    originals = []
    for role, name, text in (("product", "舊產品.txt", "額定電壓：48 V\n"), ("standard", "舊標準.txt", "額定電壓：48 V\n")):
        uploaded = ok(client.post(f"/api/projects/{project['id']}/documents", files={"file": (name, text.encode(), "text/plain")}, data={"role": role}))
        originals.append(ok(client.post(f"/api/projects/{project['id']}/documents/{uploaded['id']}/confirm")))
    before = ok(client.get(f"/api/projects/{project['id']}"))
    imported = ok(client.post(PREFIX + "/import-project", json={"project_id": project["id"]}))
    assert imported["total"] == 2 and imported["reused"] == 0
    assert {d["role"] for d in imported["items"]} == {"product", "standard"}
    assert all(not d["confirmed"] and d["item_count"] == 0 for d in imported["items"])
    assert all(d["index_status"] != "ready" for d in imported["items"])
    for new in imported["items"]:
        old = next(d for d in originals if d["role"] == new["role"])
        assert new["id"] != old["id"]
        assert client.get(f"{PREFIX}/documents/{new['id']}/original").content == client.get(f"/api/projects/{project['id']}/documents/{old['id']}/original").content
    again = ok(client.post(PREFIX + "/import-project", json={"project_id": project["id"]}))
    assert again["total"] == again["reused"] == 2
    assert {d["id"] for d in again["items"]} == {d["id"] for d in imported["items"]}
    assert ok(client.get(f"/api/projects/{project['id']}")) == before
    assert ok(client.get(PREFIX + "/library"))["total"] == ok(client.get(PREFIX + "/products"))["total"] == 1
