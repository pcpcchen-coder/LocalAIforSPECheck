"""API acceptance tests for complete, reviewable, local-only comparison runs.

No real model or network connection is required. The deterministic built-in demo
exercises the real comparison engine; local-mode tests replace only model work.
"""
from __future__ import annotations

import copy
import hashlib
import io
import threading
import time
import zipfile
from collections import Counter
from contextlib import contextmanager

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import spec_check.app as app_module
from spec_check.app import create_app
from spec_check.engine import ComparisonCancelled


BASE_URL = "http://127.0.0.1:8765"


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path), base_url=BASE_URL) as current:
        yield current


def success(response, status=200):
    assert response.status_code == status, response.text
    return response.json()


def wait_for_run(client, run_id, statuses=("completed",), timeout=8):
    deadline = time.monotonic() + timeout
    latest = None
    while time.monotonic() < deadline:
        latest = success(client.get(f"/api/runs/{run_id}"))
        if latest["status"] in statuses:
            return latest
        time.sleep(0.01)
    pytest.fail(f"Run did not reach {statuses}: {latest}")


def upload(client, project_id, name, text, role, confirm=True):
    document = success(client.post(
        f"/api/projects/{project_id}/documents",
        files={"file": (name, text.encode("utf-8"), "text/plain")},
        data={"role": role},
    ))
    if confirm:
        document = success(client.post(
            f"/api/projects/{project_id}/documents/{document['id']}/confirm"
        ))
    return document


def local_project(client, requirements="額定電壓：48 V\n防護等級：IP65\n"):
    project = success(client.post("/api/projects", json={"name": "規格驗收"}))
    product = upload(client, project["id"], "產品.txt", "額定電壓：48 V\n防護等級：IP54\n", "product")
    standard = upload(client, project["id"], "規範.txt", requirements, "standard")
    success(client.put("/api/settings", json={"model": "local-test-model"}))
    return project, product, standard


def simulated_result(requirement, product_blocks, settings, mode="local", cancel_check=None, progress_callback=None):
    """A plausible cited response, without testing model correctness here."""
    assert mode == "local"
    if cancel_check and cancel_check():
        raise ComparisonCancelled("test cancellation")
    source = product_blocks[0]
    return {
        "status": "partial", "explanation": "產品有額定值，但部分規格需要人工確認。",
        "differences": ["需要確認防護等級與測試條件。"],
        "evidence": [{"block_id": source["id"], "location": source["location"], "quote": source["text"]}],
        "confidence": 0.7, "product_coverage": {"scanned": 1, "total": 1}, "warnings": [],
    }


def review(client, row, decision="confirmed", **overrides):
    body = {"decision": decision, "reviewer": "王工程師", "note": "已核對產品原文。",
            "expected_version": row["review"]["version"]}
    body.update(overrides)
    return client.post(f"/api/results/{row['id']}/reviews", json=body)


def test_demo_every_block_review_history_rankings_and_all_exports(client):
    project = success(client.post("/api/demo"))
    run = success(client.post(f"/api/projects/{project['id']}/runs", json={"mode": "demo"}))
    run = wait_for_run(client, run["id"])
    expected = {(d["id"], b["id"]) for d in project["documents"] if d["role"] == "standard" for b in d["blocks"]}
    assert {(r["standard_id"], r["block_id"]) for r in run["results"]} == expected
    assert run["total"] == run["completed"] == len(expected) == 10
    assert {r["status"] for r in run["results"]} == {"match", "partial", "mismatch", "missing", "uncertain"}
    assert run["rankings"][0]["standard_name"] == "規範_A.txt"
    assert run["rankings"][0]["score"] == 50.0
    assert run["rankings"][1]["score"] == 37.5

    before = copy.deepcopy(run)
    for row in run["results"]:
        confirmed = success(review(client, row, final_status="match"))
        # Confirmation cannot silently replace the original AI verdict.
        assert confirmed["review"]["final_status"] == row["status"]
        assert confirmed["review"]["version"] == 1
        assert len(confirmed["history"]) == 1
    current = success(client.get(f"/api/runs/{run['id']}"))
    assert sum(r["reviewed"] for r in current["rankings"]) == len(expected)

    target = next(r for r in current["results"] if r["status"] == "mismatch")
    original = copy.deepcopy(target)
    changed = success(review(client, target, "changed", final_status="match", note="確認此要求不適用此型號，人工改判。"))
    assert changed["status"] == original["status"]
    assert changed["review"]["final_status"] == "match"
    assert [e["version"] for e in changed["history"]] == [1, 2]
    stale = review(client, target, "changed", final_status="missing", note="過期視窗送出。")
    assert stale.status_code == 409
    after_conflict = success(client.get(f"/api/runs/{run['id']}"))
    persisted = next(r for r in after_conflict["results"] if r["id"] == target["id"])
    assert persisted == changed

    reopened = success(review(client, changed, "reopened", note="二次覆核需補充適用條件。"))
    assert reopened["review"]["final_status"] is None
    assert [e["decision"] for e in reopened["history"]] == ["confirmed", "changed", "reopened"]
    final_run = success(client.get(f"/api/runs/{run['id']}"))
    assert sum(r["reviewed"] for r in final_run["rankings"]) == len(expected) - 1
    assert [r["score"] for r in final_run["rankings"]] == [r["score"] for r in before["rankings"]]
    assert final_run["documents"] == before["documents"]

    for fmt in ("json", "html", "xlsx"):
        response = client.get(f"/api/runs/{run['id']}/export", params={"format": fmt})
        assert response.status_code == 200, response.text
        assert "attachment" in response.headers["content-disposition"]
        if fmt == "json":
            exported = response.json()
            assert exported.pop("export_id")
            assert exported.pop("exported_at")
            assert exported == final_run
        elif fmt == "html":
            assert "二次覆核需補充適用條件。" in response.text
            assert "未呼叫本地模型" in response.text
            for row in final_run["results"]:
                assert row["id"] in response.text
        else:
            book = load_workbook(io.BytesIO(response.content), read_only=True)
            assert set(book.sheetnames) == {"Results", "History", "Documents", "Run"}
            assert book["Results"].max_row == len(expected) + 1
            assert book["History"].max_row == len(expected) + 3
            assert book["Documents"].max_row == 15
            ids = {values[0] for values in book["Results"].iter_rows(min_row=2, values_only=True)}
            assert ids == {r["id"] for r in final_run["results"]}
            book.close()


def test_local_run_requires_source_confirmation_and_keeps_run_snapshots(client, monkeypatch):
    calls = []

    def compare(*args, **kwargs):
        calls.append((args, kwargs))
        return simulated_result(*args, **kwargs)

    monkeypatch.setattr(app_module, "compare_block", compare)
    project = success(client.post("/api/projects", json={"name": "同型號版本追蹤"}))
    product = upload(client, project["id"], "產品.txt", "額定值：48 V\n", "product", confirm=False)
    standard = upload(client, project["id"], "原規範.txt", "額定值：48 V\n", "standard")
    endpoint = f"/api/projects/{project['id']}/runs"
    success(client.put("/api/settings", json={"model": "model-v1"}))
    assert client.post(endpoint, json={}).status_code == 400
    success(client.post(f"/api/projects/{project['id']}/documents/{product['id']}/confirm"))
    run = success(client.post(endpoint, json={"standard_ids": [standard["id"]]}))
    run = wait_for_run(client, run["id"])
    assert len(calls) == len(standard["blocks"])
    assert run["mode"] == "local"
    assert calls[0][0][2]["model"] == "model-v1"
    original = copy.deepcopy(run)

    success(client.delete(f"/api/projects/{project['id']}/documents/{product['id']}"))
    replacement = upload(client, project["id"], "產品新版.txt", "額定值：24 V\n", "product")
    new_standard = upload(client, project["id"], "新規範.txt", "通訊：CAN\n", "standard")
    success(client.put("/api/settings", json={"model": "model-v2"}))
    assert success(client.get(f"/api/runs/{run['id']}")) == original
    assert client.get(f"/api/projects/{project['id']}/documents/{product['id']}/original").content == "額定值：48 V\n".encode()
    next_run = success(client.post(endpoint, json={"standard_ids": [new_standard["id"]]}))
    next_run = wait_for_run(client, next_run["id"])
    assert next_run["id"] != run["id"]
    assert set(next_run["document_ids"]) == {replacement["id"], new_standard["id"]}
    assert next_run["settings"]["model"] == "model-v2"
    assert next_run["results"][0]["evidence"][0]["quote"] == "額定值：24 V"


def test_model_failure_keeps_every_requirement_and_never_leaks_secret(client, monkeypatch):
    secret = "test-local-secret-924caf"
    project, product, standard = local_project(client)
    settings = success(client.put("/api/settings", json={"api_key": secret}))
    assert settings["has_api_key"] is True
    assert "api_key" not in settings
    received = []

    def failed_model(requirement, product_blocks, settings, **kwargs):
        received.append(settings["api_key"])
        raise RuntimeError("server error echoing Authorization: " + settings["api_key"])

    monkeypatch.setattr(app_module, "compare_block", failed_model)
    run = success(client.post(f"/api/projects/{project['id']}/runs", json={}))
    run = wait_for_run(client, run["id"])
    assert received == [secret] * len(standard["blocks"])
    assert run["completed"] == run["total"] == len(standard["blocks"])
    assert all(r["status"] == "uncertain" and not r["evidence"] for r in run["results"])
    assert all(r["product_coverage"]["scanned"] == 0 for r in run["results"])
    assert run["rankings"][0]["score"] == 0

    for endpoint in ("/api/settings", "/api/projects", f"/api/projects/{project['id']}",
                     f"/api/projects/{project['id']}/runs", f"/api/runs/{run['id']}"):
        response = client.get(endpoint)
        assert response.status_code == 200
        assert secret not in response.text
    for fmt in ("html", "json", "xlsx"):
        response = client.get(f"/api/runs/{run['id']}/export", params={"format": fmt})
        assert response.status_code == 200
        if fmt == "xlsx":
            with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                assert all(secret.encode() not in archive.read(name) for name in archive.namelist())
        else:
            assert secret not in response.text

    # Empty UI password inputs preserve the stored token; explicit clear erases it.
    assert success(client.put("/api/settings", json={"api_key": "", "temperature": 0.2}))["has_api_key"]
    assert not success(client.put("/api/settings", json={"clear_api_key": True}))["has_api_key"]
    assert not success(client.get("/api/settings"))["has_api_key"]


def test_archived_reports_preserve_bytes_reviews_hashes_and_hide_secrets(client):
    secret = "archived-local-key-24ab9d"
    success(client.put("/api/settings", json={"api_key": secret}))
    project = success(client.post("/api/demo"))
    run = success(client.post(f"/api/projects/{project['id']}/runs", json={"mode": "demo"}))
    run = wait_for_run(client, run["id"])
    history_url = f"/api/runs/{run['id']}/exports"
    assert success(client.get(history_url)) == []
    archives = []
    for fmt in ("json", "html", "xlsx"):
        response = client.get(f"/api/runs/{run['id']}/export", params={"format": fmt})
        assert response.status_code == 200, response.text
        history = success(client.get(history_url))
        item = history[-1]
        assert len(history) == len(archives) + 1
        assert item["run_id"] == run["id"]
        assert item["format"] == fmt
        assert item["filename"].endswith("." + fmt)
        assert item["created_at"]
        assert item["sha256"] == hashlib.sha256(response.content).hexdigest()
        assert not any(key.startswith("_") for key in item)
        downloaded = client.get(f"/api/exports/{item['id']}/download")
        assert downloaded.status_code == 200
        assert downloaded.content == response.content
        if fmt == "json":
            snapshot = downloaded.json()
            assert snapshot["export_id"] == item["id"]
            assert snapshot["exported_at"] == item["created_at"]
            assert all(r["review"]["version"] == 0 and r["history"] == [] for r in snapshot["results"])
        if fmt == "xlsx":
            with zipfile.ZipFile(io.BytesIO(downloaded.content)) as zipped:
                assert all(secret.encode() not in zipped.read(name) for name in zipped.namelist())
        else:
            assert secret not in downloaded.text
        archives.append((item, response.content))

    # The current workspace moves forward while saved review artifacts remain exact.
    changed = success(review(client, run["results"][0], "changed", final_status="uncertain", note="補上新確認過程，舊報告須保留原樣。"))
    current = success(client.get(f"/api/runs/{run['id']}"))
    assert current["exports"] == [item for item, _ in archives]
    assert next(r for r in current["results"] if r["id"] == changed["id"]) == changed
    for item, content in archives:
        downloaded = client.get(f"/api/exports/{item['id']}/download")
        assert downloaded.status_code == 200
        assert downloaded.content == content
    newer = client.get(f"/api/runs/{run['id']}/export", params={"format": "json"})
    assert newer.status_code == 200
    newer_snapshot = newer.json()
    newer_row = next(r for r in newer_snapshot["results"] if r["id"] == changed["id"])
    assert newer_row["review"]["version"] == newer_row["history"][-1]["version"] == 1
    assert newer_row["history"][0]["note"] == changed["review"]["note"]
    history = success(client.get(history_url))
    assert history[:3] == [item for item, _ in archives]
    assert len(history) == 4
    assert len({item["id"] for item in history}) == 4
    assert secret not in client.get(history_url).text

    # External tampering must not be delivered as if it were the saved report.
    store = client.app.state.store
    corrupted = store.get("export", archives[0][0]["id"])
    path = store.root / "reports" / corrupted["_stored_name"]
    path.write_bytes(b"externally replaced content")
    assert client.get(f"/api/exports/{corrupted['id']}/download").status_code == 409
    path.unlink()
    assert client.get(f"/api/exports/{corrupted['id']}/download").status_code == 404
    assert success(client.get(history_url)) == history


def test_run_snapshot_keeps_review_and_history_at_one_version_during_concurrent_write(client, monkeypatch):
    project = success(client.post("/api/demo"))
    run = success(client.post(f"/api/projects/{project['id']}/runs", json={"mode": "demo"}))
    run = wait_for_run(client, run["id"])
    target = run["results"][0]
    store = client.app.state.store
    connect = store.connect
    committed = []
    failures = []
    triggered = threading.Event()

    def write_review():
        try:
            committed.append(store.review(target["id"], {
                "decision": "confirmed", "reviewer": "同時覆核者", "note": "快照讀取期間提交。",
                "expected_version": 0,
            }))
        except BaseException as exc:
            failures.append(exc)

    def between_read_statements(statement):
        # A trace hook pauses exactly before the audit-history SELECT. The writer
        # commits on its own connection while this reader's transaction is open.
        # Without a stable snapshot this would pair verdict v0 with history v1.
        if "FROM reviews" not in statement or triggered.is_set():
            return
        triggered.set()
        writer = threading.Thread(target=write_review, name="concurrent-review")
        writer.start()
        writer.join(5)
        if writer.is_alive():
            failures.append(AssertionError("Concurrent writer did not finish"))

    @contextmanager
    def connect_with_read_barrier():
        with connect() as db:
            if threading.current_thread().name != "concurrent-review":
                db.set_trace_callback(between_read_statements)
            yield db

    monkeypatch.setattr(store, "connect", connect_with_read_barrier)
    snapshot = store.run_snapshot(run["id"])
    assert triggered.is_set(), "The test did not overlap the read with a review commit"
    assert not failures
    assert committed[0]["review"]["version"] == 1
    observed = next(r for r in snapshot["results"] if r["id"] == target["id"])
    assert observed["review"]["version"] == 0
    assert observed["history"] == []
    # The next independent snapshot sees both parts of the committed transaction.
    fresh = store.run_snapshot(run["id"])
    observed = next(r for r in fresh["results"] if r["id"] == target["id"])
    assert observed["review"]["version"] == observed["history"][-1]["version"] == 1


def test_connection_uses_saved_key_without_returning_or_overwriting_it(client, monkeypatch):
    secret = "connection-token-ef981"
    success(client.put("/api/settings", json={"api_key": secret, "model": "saved-model"}))
    received = []

    def connected(settings):
        received.append(settings)
        return [{"id": "available-model"}]

    monkeypatch.setattr(app_module, "check_connection", connected)
    response = success(client.post("/api/connection", json={"settings": {"model": "temporary-model", "api_key": ""}}))
    assert response["models"] == [{"id": "available-model"}]
    assert received[0]["api_key"] == secret
    assert received[0]["model"] == "temporary-model"
    assert success(client.get("/api/settings"))["model"] == "saved-model"

    def failed(settings):
        raise RuntimeError(secret)

    monkeypatch.setattr(app_module, "check_connection", failed)
    response = client.post("/api/connection", json={})
    assert response.status_code == 502
    assert secret not in response.text


@pytest.mark.parametrize("role,name", [("unknown", "file.txt"), ("standard", "old.doc"), ("standard", "program.exe")])
def test_invalid_upload_does_not_create_a_document_or_leave_files(client, role, name):
    project = success(client.post("/api/projects", json={"name": "上傳限制"}))
    response = client.post(f"/api/projects/{project['id']}/documents",
                           files={"file": (name, b"test")}, data={"role": role})
    assert response.status_code == 400
    assert not success(client.get(f"/api/projects/{project['id']}"))["documents"]
    assert not list((client.app.state.store.root / "uploads").iterdir())


def test_duplicate_product_unknown_standard_and_demo_misuse_are_rejected(client):
    project, product, standard = local_project(client)
    root = client.app.state.store.root
    previous_files = set((root / "uploads").iterdir())
    response = client.post(f"/api/projects/{project['id']}/documents", files={"file": ("duplicate.txt", b"48 V")}, data={"role": "product"})
    assert response.status_code == 400
    assert set((root / "uploads").iterdir()) == previous_files
    assert len(success(client.get(f"/api/projects/{project['id']}"))["documents"]) == 2
    endpoint = f"/api/projects/{project['id']}/runs"
    assert client.post(endpoint, json={"mode": "demo"}).status_code == 400
    assert client.post(endpoint, json={"standard_ids": ["nonexistent"]}).status_code == 400
    assert not success(client.get(endpoint))


def test_review_validation_does_not_change_audit_history(client):
    project = success(client.post("/api/demo"))
    run = success(client.post(f"/api/projects/{project['id']}/runs", json={"mode": "demo"}))
    run = wait_for_run(client, run["id"])
    row = run["results"][0]
    assert review(client, row, "changed", final_status="match", note="  ").status_code == 400
    assert review(client, row, "changed", note="有原因但沒有判定").status_code == 400
    assert review(client, row, "reopened", note="").status_code == 400
    assert review(client, row, reviewer=" ").status_code == 400
    assert review(client, row, expected_version=-1).status_code == 422
    unchanged = success(client.get(f"/api/runs/{run['id']}"))
    assert unchanged == run


def test_cancel_and_resume_keep_completed_rows_and_original_model_settings(client, monkeypatch):
    project, product, standard = local_project(client, "條款一：48 V\n條款二：IP65\n條款三：CAN\n")
    entered_second = threading.Event()
    release_second = threading.Event()
    calls = []

    def blocking(requirement, product_blocks, settings, mode="local", cancel_check=None, progress_callback=None):
        calls.append((requirement["id"], settings["model"]))
        if len(calls) == 2:
            entered_second.set()
            assert release_second.wait(5), "test did not release model response"
        return simulated_result(requirement, product_blocks, settings, mode, cancel_check)

    monkeypatch.setattr(app_module, "compare_block", blocking)
    run = success(client.post(f"/api/projects/{project['id']}/runs", json={}))
    try:
        assert entered_second.wait(5)
        in_progress = success(client.get(f"/api/runs/{run['id']}"))
        assert in_progress["completed"] == 1
        assert in_progress["rankings"][0]["total"] == 3
        first = in_progress["results"][0]
        first = success(review(client, first))
        success(client.post(f"/api/runs/{run['id']}/cancel"))
    finally:
        release_second.set()
    cancelled = wait_for_run(client, run["id"], ("cancelled",))
    assert cancelled["completed"] == 1
    assert cancelled["results"][0] == first
    success(client.put("/api/settings", json={"model": "new-model-for-future-runs"}))
    success(client.post(f"/api/runs/{run['id']}/resume"))
    finished = wait_for_run(client, run["id"])
    assert finished["completed"] == finished["total"] == 3
    assert finished["results"][0] == first
    assert len({r["id"] for r in finished["results"]}) == 3
    assert Counter(block_id for block_id, _ in calls) == {"B00001": 1, "B00002": 2, "B00003": 1}
    assert {model for _, model in calls} == {"local-test-model"}
    assert client.post(f"/api/runs/{run['id']}/resume").status_code == 400


def test_restart_marks_unfinished_runs_interrupted_and_resumes_remaining_work(tmp_path, monkeypatch):
    application = create_app(tmp_path)
    with TestClient(application, base_url=BASE_URL) as client:
        project = success(client.post("/api/demo"))
        run = success(client.post(f"/api/projects/{project['id']}/runs", json={"mode": "demo"}))
        run = wait_for_run(client, run["id"])
        preserved = success(review(client, run["results"][0]))
    # Reproduce durable state left by a process killed after saving all but one row.
    store = application.state.store
    last = run["results"][-1]
    with store.connect() as db:
        db.execute("DELETE FROM records WHERE kind='result' AND id=?", (last["id"],))
    stored = store.get("run", run["id"])
    stored.update(status="running", finished_at=None)
    store.put("run", stored, project["id"])
    original_compare = app_module.compare_block
    resumed_calls = []

    def recording(requirement, *args, **kwargs):
        resumed_calls.append(requirement["text"])
        return original_compare(requirement, *args, **kwargs)

    monkeypatch.setattr(app_module, "compare_block", recording)
    with TestClient(create_app(tmp_path), base_url=BASE_URL) as restarted:
        interrupted = success(restarted.get(f"/api/runs/{run['id']}"))
        assert interrupted["status"] == "interrupted"
        assert interrupted["completed"] == run["total"] - 1
        success(restarted.post(f"/api/runs/{run['id']}/resume"))
        completed = wait_for_run(restarted, run["id"])
        assert completed["completed"] == run["total"]
        assert resumed_calls == [last["requirement"]]
        assert next(r for r in completed["results"] if r["id"] == preserved["id"]) == preserved


def test_local_host_origin_and_metadata_safety_headers(client):
    health = client.get("/api/health")
    assert health.status_code == 200
    assert health.json()["status"] == "ok"
    assert health.headers["cache-control"] == "no-store"
    assert health.headers["x-content-type-options"] == "nosniff"
    assert "frame-ancestors 'none'" in health.headers["content-security-policy"]
    assert client.get("/api/health", headers={"host": "attacker.example"}).status_code == 403
    assert client.post("/api/projects", json={"name": "cross-origin"}, headers={"origin": "https://attacker.example"}).status_code == 403
    assert client.post("/api/projects", json={"name": "cross-site"}, headers={"sec-fetch-site": "cross-site"}).status_code == 403
    assert not success(client.get("/api/projects"))
    assert client.post("/api/projects", json={"name": "same-origin"}, headers={"origin": BASE_URL}).status_code == 200
    assert client.put("/api/settings", json={"base_url": "https://api.example.com/v1"}).status_code == 400
    assert client.get("/api/runs/no-such-run").status_code == 404
