import copy
import io
import json
import re

import pytest
from openpyxl import load_workbook

from spec_check.exports import export_run


@pytest.fixture
def analysis():
    product_item = {"id": "p1", "kind": "specification", "name": "通訊", "quote": "支援 CAN", "block_id": "B1"}
    voltage = {"id": "s1", "kind": "requirement", "name": "電壓", "value": "48", "unit": "V", "quote": "電壓 48 V", "block_id": "B1"}
    protection = {"id": "s2", "kind": "requirement", "name": "防水", "quote": "IP67", "block_id": "B2"}
    unfinished = {"id": "s3", "kind": "unresolved", "name": "待解析試驗", "quote": "尚未完成的要求", "block_id": "B3"}
    return {
        "schema_version": 2, "kind": "analysis", "id": "a1", "name": "產品初審", "status": "cancelled", "comparison_mode": "focused",
        "settings": {"api_key": "DO-NOT-EXPORT"},
        "coverage": {"screened": 2, "total_standards": 2, "selected_standards": 1, "excluded_standards": 1, "compared": 2, "total_items": 3, "coverage_complete": False},
        "documents": [
            {"id": "product", "name": "產品.txt", "role": "product", "items": [product_item], "blocks": [{"id": "B1", "text": "支援 CAN", "location": "第 1 行"}], "extraction_audit": [{"id": "extract-product", "action": "split", "note": "產品修訂歷程"}]},
            {"id": "standard", "name": "電源標準", "role": "standard", "items": [voltage, protection, unfinished], "blocks": [{"id": "B1", "text": "電壓 48 V"}, {"id": "B2", "text": "IP67"}, {"id": "B3", "text": "尚未完成的要求"}], "extraction_audit": [{"id": "extract-standard", "action": "item_update", "note": "校對例外條件"}]},
            {"id": "excluded", "name": "運輸標準", "role": "standard", "items": [{"id": "excluded-item", "kind": "requirement", "name": "運輸要求", "quote": "運輸"}], "blocks": [{"id": "B1", "text": "運輸"}], "extraction_audit": []},
        ],
        "screenings": [
            {"id": "screen-1", "standard_id": "standard", "standard_name": "電源標準", "relevance": "high", "applicability": "unknown", "reason": "產品含有電源", "selected": True, "review": {"reviewer": "George", "note": "需要檢查"}},
            {"id": "screen-2", "standard_id": "excluded", "standard_name": "運輸標準", "relevance": "unknown", "applicability": "unknown", "reason": "運送條件未知", "selected": False, "review": {"reviewer": "George", "reason": "本批不含運送範圍", "note": "人工排除理由"}},
        ],
        "results": [
            {"id": "r1", "standard_id": "standard", "standard_name": "電源標準", "item": voltage, "requirement": "電壓 48 V", "status": "mismatch", "risk": {"level": "high", "kind": "nonconformity", "basis": "電壓不同", "action": "確認規格"}, "ai_risk": {"level": "high", "kind": "nonconformity", "basis": "電壓不同", "action": "確認規格"}, "evidence": [{"block_id": "B2", "quote": "額定 48.0 V"}], "matched_product_item_ids": [], "retrieval": {"mode": "focused", "retrieved": 1, "total": 2}, "review": {"decision": "changed", "final_status": "partial", "risk_level": "low", "reviewer": "George", "note": "已核對容差", "version": 2}, "history": [{"id": "review-before", "decision": "confirmed", "risk_level": "high", "version": 1}, {"id": "review-after", "decision": "changed", "risk_level": "low", "note": "已核對容差", "version": 2}]},
            {"id": "r2", "standard_id": "standard", "standard_name": "電源標準", "item": protection, "requirement": "IP67", "status": "missing", "risk": {"level": "unknown", "kind": "evidence_gap", "basis": "沒有防水證據", "action": "補防水試驗報告"}, "ai_risk": {"level": "unknown", "kind": "evidence_gap"}, "evidence": [], "review": {"decision": "pending", "version": 0}, "history": []},
        ],
        "product_extras": {"items": [product_item], "coverage_complete": False, "note": "分析尚未完成"},
        "audit": [{"id": "selection-1", "action": "selection", "decisions": [{"standard_id": "excluded", "selected": False, "reason": "本批不含運送範圍"}]}, {"id": "analysis-created", "action": "created"}],
        "rankings": [],
    }


def _section(text, identifier):
    return text.split(f'id="{identifier}">', 1)[1].split("<h2", 1)[0]


def _rows(sheet):
    rows = list(sheet.values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def test_analysis_html_exposes_exclusions_effective_risks_and_partial_work(analysis):
    data, _, _ = export_run(analysis, "html")
    text = data.decode()
    assert "部分分析" in text
    assert "結構覆蓋" in text and "語意完整性" in text
    assert "尚未完成的要求" in _section(text, "coverage")
    assert "運輸要求" not in _section(text, "coverage")
    assert "人工排除理由" in _section(text, "screenings")
    assert "不代表已證明不適用" in _section(text, "screenings")
    risks = _section(text, "risks")
    assert re.search(r"部分符合</pre></td><td><pre>低</pre></td><td><pre>高", risks)
    assert "已核對容差" in risks
    assert "補防水試驗報告" in _section(text, "gaps")
    assert "額定 48.0 V" in _section(text, "results")
    assert "未對應不代表標準沒有要求" in _section(text, "extras")
    assert "review-before" in text and "review-after" in text
    assert "extract-product" in text and "extract-standard" in text
    assert "selection-1" in _section(text, "audit")
    assert "DO-NOT-EXPORT" not in text
    assert "分數＝" not in text


def test_analysis_workbook_has_actionable_sheets_and_all_audits(analysis):
    data, _, _ = export_run(analysis, "xlsx")
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    assert {"Results", "History", "Documents", "Run", "Screenings", "Risks", "EvidenceGaps", "ProductExtras", "Items", "SelectionAudit", "ExtractionAudit", "Audit", "Unprocessed"} == set(workbook.sheetnames)
    risk = next(row for row in _rows(workbook["Risks"]) if row["結果 ID"] == "r1")
    assert (risk["目前風險"], risk["AI 風險"], risk["目前判定"]) == ("low", "high", "partial")
    assert risk["覆核原因"] == "已核對容差"
    result = _rows(workbook["Results"])[0]
    assert result["目前風險"] == "low" and result["項目 ID"] == "s1"
    assert [row["事件 ID"] for row in _rows(workbook["History"])] == ["review-before", "review-after"]
    assert _rows(workbook["Unprocessed"])[0]["項目 ID"] == "s3"
    assert len(_rows(workbook["Unprocessed"])) == 1
    assert len(_rows(workbook["Items"])) == 5
    assert len(_rows(workbook["ExtractionAudit"])) == 2
    assert len(_rows(workbook["Audit"])) == 2
    assert "selection-1" in _rows(workbook["SelectionAudit"])[-1]["完整選取事件 JSON"]
    assert _rows(workbook["ProductExtras"])[0]["完整覆蓋"] == "False"
    assert _rows(workbook["EvidenceGaps"])[0]["建議動作"] == "補防水試驗報告"
    assert dict((row["欄位"], row["值"]) for row in _rows(workbook["Run"]))["partial_analysis"] == "True"


def test_analysis_all_formats_escape_sources_and_preserve_long_audit(analysis):
    attack = '=HYPERLINK("https://example.invalid", "click")'
    analysis["screenings"][0]["standard_name"] = '<img src=x onerror="alert(1)">'
    analysis["screenings"][0]["reason"] = attack
    long_action = "證據" * 36_000
    analysis["results"][0]["risk"]["action"] = long_action
    analysis["audit"][0]["secret"] = "HIDDEN-SECRET"
    before = copy.deepcopy(analysis)
    html, _, _ = export_run(analysis, "html")
    assert b"<img " not in html
    assert b"&lt;img" in html and b"HIDDEN-SECRET" not in html
    data, _, _ = export_run(analysis, "xlsx")
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    assert workbook["Screenings"]["F2"].value == attack
    assert workbook["Screenings"]["F2"].data_type == "s"
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    risk_rows = _rows(workbook["Risks"])
    first = next(index for index, row in enumerate(risk_rows) if row["結果 ID"] == "r1")
    assert "".join(row["建議動作"] or "" for row in risk_rows[first:]) == long_action
    raw, _, _ = export_run(analysis, "json")
    report = json.loads(raw)
    assert report["results"][0]["risk"]["action"] == long_action
    assert report["documents"] == analysis["documents"]
    assert b"HIDDEN-SECRET" not in raw
    assert analysis == before


def test_reopened_review_does_not_override_current_ai_risk(analysis):
    analysis["results"][0]["review"]["decision"] = "reopened"
    data, _, _ = export_run(analysis, "xlsx")
    workbook = load_workbook(io.BytesIO(data))
    risk = next(row for row in _rows(workbook["Risks"]) if row["結果 ID"] == "r1")
    assert risk["目前風險"] == "high" and risk["目前判定"] == "mismatch"


def test_terminal_status_alone_never_claims_complete_analysis(analysis):
    analysis["status"] = "completed"
    analysis["coverage"]["coverage_complete"] = True
    text = export_run(analysis, "html")[0].decode()
    assert "部分分析" in text  # A selected source requirement has no result.
    analysis["documents"][1]["items"].pop()
    text = export_run(analysis, "html")[0].decode()
    assert "本報告為部分分析" not in text
    assert "也不表示全部條件、例外與差異均已辨識" in text
