"""Self-contained review exports with complete source snapshots and no API keys."""

from __future__ import annotations

import html
import io
import json
import re
from datetime import datetime, timezone

STATUSES = {"match": "符合", "partial": "部分符合", "mismatch": "不符", "missing": "產品文件未載明", "uncertain": "待釐清"}
DECISIONS = {"pending": "待覆核", "confirmed": "確認 AI 判定", "changed": "人工更正", "reopened": "重新開啟"}
_SECRET_KEYS = {"api_key", "apikey", "authorization", "access_token", "refresh_token", "password", "secret", "client_secret"}


def _sanitize(value):
    if isinstance(value, dict):
        return {str(key): _sanitize(item) for key, item in value.items() if str(key).lower().replace("-", "_") not in _SECRET_KEYS}
    if isinstance(value, (list, tuple)):
        return [_sanitize(item) for item in value]
    return value


def _json(value) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2, default=str)


def _escape(value) -> str:
    return html.escape(str(value if value is not None else ""), quote=True)


def _effective(row: dict) -> str:
    review = row.get("review") or {}
    if review.get("decision") in {"confirmed", "changed"} and review.get("final_status"):
        return review["final_status"]
    return row.get("status", "uncertain")


def _paragraph(value, css="") -> str:
    return f'<pre class="{css}">{_escape(value)}</pre>'


def _list(values) -> str:
    return "<ul>" + "".join(f"<li>{_escape(value)}</li>" for value in values) + "</ul>" if values else "<p>無</p>"


RISK_LEVELS = {"high": "高", "medium": "中", "low": "低", "unknown": "待確認", "none": "無已辨識提示"}
RISK_KINDS = {"nonconformity": "規格差異", "evidence_gap": "證據缺漏", "uncertainty": "判定不確定", "none": "無"}
RELEVANCE = {"high": "高度相關", "medium": "可能相關", "low": "低相關", "unknown": "待確認", "related": "相關", "unrelated": "未見相關"}
APPLICABILITY = {"likely": "可能適用", "conditional": "有條件適用", "applicable": "適用", "potentially_applicable": "可能適用", "not_applicable": "不適用", "unknown": "待確認", "uncertain": "待確認"}


def _is_analysis(run: dict) -> bool:
    return run.get("kind") == "analysis"


def _risk_value(value) -> dict:
    if isinstance(value, dict):
        return dict(value)
    return {"level": value or "unknown"}


def _effective_risk(row: dict) -> dict:
    risk = _risk_value(row.get("risk") or row.get("ai_risk"))
    review = row.get("review") or {}
    if review.get("decision") in {"confirmed", "changed"} and review.get("risk_level"):
        risk["level"] = review["risk_level"]
        risk["reviewed"] = True
        if review.get("note"):
            risk["review_note"] = review["note"]
    return risk


def _risk_label(value) -> str:
    return RISK_LEVELS.get(value, str(value or "待確認"))


def _item_id(row: dict) -> str:
    item = row.get("item") or {}
    return str(row.get("item_id") or (item.get("id") if isinstance(item, dict) else "") or "")


def _extra_items(run: dict) -> tuple[list, dict]:
    extras = run.get("product_extras") or {}
    if isinstance(extras, list):
        return extras, {}
    return extras.get("items") or [], extras


def _unprocessed_items(run: dict) -> list[dict]:
    # Show absent rows, never manufacture an AI verdict for unfinished work.
    selected = {row.get("standard_id") for row in run.get("screenings", []) if row.get("selected") is True}
    processed = {(row.get("standard_id"), _item_id(row)) for row in run.get("results", [])}
    missing = []
    for document in run.get("documents", []):
        if document.get("id") not in selected:
            continue
        for item in document.get("items", []):
            if item.get("kind") not in {"requirement", "unresolved"}:
                continue
            if (document.get("id"), str(item.get("id", ""))) not in processed:
                missing.append({"standard_id": document.get("id"), "standard_name": document.get("name"), "item": item})
    return missing


def _analysis_partial(run: dict) -> bool:
    coverage = run.get("coverage") or {}
    return run.get("status") != "completed" or coverage.get("coverage_complete") is not True or bool(_unprocessed_items(run))


def _priority_rows(run: dict) -> list[dict]:
    order = {"high": 0, "medium": 1, "unknown": 2, "low": 3, "none": 4}
    rows = [row for row in run.get("results", []) if not row.get("excluded_from_requirements") and (row.get("item") or {}).get("kind") != "context" and (_effective_risk(row).get("level", "unknown") != "none" or _effective(row) != "match")]
    return sorted(rows, key=lambda row: order.get(_effective_risk(row).get("level"), 2))


def _evidence_gaps(run: dict) -> list[dict]:
    return [row for row in run.get("results", []) if not row.get("excluded_from_requirements") and (row.get("item") or {}).get("kind") != "context" and (_effective(row) in {"missing", "uncertain"} or _effective_risk(row).get("kind") in {"evidence_gap", "uncertainty"})]


def _evidence_text(evidence: list) -> str:
    return "\n\n".join(f"{entry.get('document_id', '')} {entry.get('block_id', '')} {entry.get('location', '')}\n{entry.get('quote', '')}".strip() for entry in evidence) or "尚無有效原文引述"


def _risk_summary(row: dict) -> str:
    risk = _effective_risk(row)
    summary = f"類型：{RISK_KINDS.get(risk.get('kind'), risk.get('kind', '待確認'))}\n依據：{risk.get('basis', '')}\n建議動作：{risk.get('action', '')}"
    if risk.get("review_note"):
        summary += f"\n人工覆核理由：{risk['review_note']}"
    return summary


def _item_summary(item: dict) -> str:
    fields = [("parameter", "參數"), ("value", "數值"), ("unit", "單位"), ("operator", "比較關係"), ("conditions", "適用條件"), ("exceptions", "例外"), ("test_method", "測試方法"), ("criticality", "重要性"), ("criticality_basis", "重要性依據")]
    return "\n".join(f"{label}：{_risk_label(item[key]) if key == 'criticality' else item[key]}" for key, label in fields if item.get(key)) or "尚未抽取完整規格欄位，請核對原文。"


def _selection_text(entry: dict) -> str:
    selected = "選取" if entry.get("selected") is True else "人工排除／未選取" if entry.get("selected") is False else "選取狀態待確認"
    review = entry.get("review") or {}
    return f"{selected}\n覆核人員：{review.get('reviewer', '')}\n理由：{review.get('reason') or review.get('note', '')}\n註記：{review.get('note', '')}"


def _table(headings: list, rows: list) -> str:
    head = "".join(f"<th>{_escape(value)}</th>" for value in headings)
    body = "".join("<tr>" + "".join(f"<td><pre>{_escape(_json(value) if isinstance(value, (dict, list, tuple)) else value)}</pre></td>" for value in row) + "</tr>" for row in rows)
    if not rows:
        body = f'<tr><td colspan="{len(headings)}">目前沒有已產生的項目；請一併查看處理覆蓋與未處理清單。</td></tr>'
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def _analysis_html(run: dict) -> bytes:
    parts = ["""<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>標準相關性與風險覆核報告</title>
<style>body{font-family:system-ui,'Microsoft JhengHei',sans-serif;color:#172b3a;max-width:1200px;margin:32px auto;padding:0 24px;line-height:1.65}h1,h2,h3{line-height:1.3}h2{border-bottom:2px solid #17696a;padding-bottom:8px;margin-top:40px}pre{font-family:inherit;white-space:pre-wrap;overflow-wrap:anywhere;margin:8px 0}table{border-collapse:collapse;width:100%;margin:12px 0}td,th{border:1px solid #c8d5dc;padding:8px;text-align:left;vertical-align:top;overflow-wrap:anywhere}th{background:#edf4f6}.row,.source{border:1px solid #c8d5dc;border-radius:8px;padding:18px;margin:16px 0}.muted{color:#526571;font-size:.9em}.notice{background:#fff4d5;border-left:4px solid #ad7700;padding:14px}.badge{display:inline-block;background:#e8f3f2;padding:3px 12px;border-radius:12px}.evidence{border-left:3px solid #408787;padding-left:14px}.metadata{font-size:.85em}a{color:#136e70}@media print{body{max-width:none;margin:0;padding:0;font-size:10pt}h2,h3{break-after:avoid}thead{display:table-header-group}.row{break-inside:auto}@page{margin:16mm}}</style></head><body>"""]
    parts.append("<h1>標準相關性與風險覆核報告</h1>" + _paragraph(f"分析：{run.get('name') or run.get('project_name', '')}\n分析 ID：{run.get('id', '')}　狀態：{run.get('status', '')}\n產品：{run.get('product_name', '')}　比對模式：{run.get('comparison_mode', '')}\n建立：{run.get('created_at', '')}　完成：{run.get('finished_at') or '尚未完成'}"))
    notice = "相關性、適用性、符合判定與風險提示分別呈現；不以符合度分數代替相關性。風險等級是覆核優先提示，並非經校準的事故機率、損失預測或法規認證。產品文件未載明不等於產品本身不合規。覆核人員為自行填寫，非驗證身分或電子簽章。"
    if _analysis_partial(run):
        notice += " 本報告為部分分析：尚未完整處理，或未涵蓋全部產品分段；未處理、未檢索及未確認項目不可視為符合或無風險。"
    parts.append('<p class="notice">' + _escape(notice) + "</p>")
    parts.append('<nav><a href="#screenings">相關標準</a> · <a href="#risks">優先覆核</a> · <a href="#gaps">待補證據</a> · <a href="#results">逐項差異</a> · <a href="#extras">未對應產品</a> · <a href="#sources">原文與抽取</a> · <a href="#audit">完整歷程</a></nav>')
    parts.append('<h2 id="coverage">處理覆蓋與分析限制</h2>')
    parts.append("<p>結構覆蓋是已處理的文件、原文區塊、項目或檢索分段數；語意完整性尚未得到保證。即使全部項目已處理，也不表示全部條件、例外與差異均已辨識。聚焦比對僅檢查取回的候選原文，未檢索部分須另行確認。</p>")
    coverage = run.get("coverage") or {}
    parts.append(_table(["處理範圍", "進度／數量"], [
        ["已初篩標準／本次標準", f"{coverage.get('screened', 0)} / {coverage.get('total_standards', 0)}"],
        ["選取／排除標準", f"{coverage.get('selected_standards', 0)} / {coverage.get('excluded_standards', 0)}"],
        ["已比對項目／本次要求項目", f"{coverage.get('compared', 0)} / {coverage.get('total_items', 0)}"],
        ["已完成全部產品原文比對的項目", coverage.get("exhaustive_results", "待確認")],
        ["尚未解析項目", coverage.get("unresolved_items", "待確認")],
        ["分析覆蓋狀態", "部分分析，仍需確認" if _analysis_partial(run) else "本次選取要求與產品原文已處理；語意仍需覆核"],
    ]))
    parts.append(_paragraph(_json({"coverage": coverage, "counts": run.get("counts", {}), "partial_analysis": _analysis_partial(run)}), "metadata"))
    if run.get("error"):
        parts.append("<h3>執行診斷</h3>" + _paragraph(run["error"]))
    unprocessed = _unprocessed_items(run)
    parts.append("<h3>尚未產生結果的已選規範項目</h3>" + _table(["規範", "項目 ID", "項目", "原文", "處理狀態"], [[entry.get("standard_name"), entry["item"].get("id"), entry["item"].get("name"), entry["item"].get("quote"), "未處理／無結果"] for entry in unprocessed]))
    parts.append('<h2 id="screenings">相關標準與適用性初篩</h2><p>本表保留本次初篩的全部標準，包含待確認及人工排除。人工排除只表示本次不進行項目比對，不代表已證明不適用。</p>')
    parts.append(_table(["標準／ID", "相關性", "適用性", "初篩理由與證據", "本次選取／人工紀錄", "警告"], [[f"{entry.get('standard_name', '')}\n{entry.get('standard_id', '')}", RELEVANCE.get(entry.get("relevance"), entry.get("relevance", "待確認")), APPLICABILITY.get(entry.get("applicability"), entry.get("applicability", "待確認")), f"理由：{entry.get('reason', '')}\n\n證據：{_evidence_text(entry.get('evidence') or [])}", _selection_text(entry), "\n".join(entry.get("warnings") or [])] for entry in run.get("screenings", [])]))
    parts.append('<h2 id="risks">優先覆核與風險提示</h2><p>依目前風險等級排列；人工確認或更正的風險覆核優先採用，原始 AI 提示另行保留。待確認不等於低風險。</p>')
    parts.append(_table(["結果／規範／項目", "目前判定", "目前風險", "原始 AI 風險", "依據與建議動作"], [[f"{row.get('id', '')}\n{row.get('standard_name', '')}\n{(row.get('item') or {}).get('name') or row.get('requirement', '')}", STATUSES.get(_effective(row), _effective(row)), _risk_label(_effective_risk(row).get("level")), _risk_label(_risk_value(row.get("ai_risk") or row.get("risk")).get("level")), _risk_summary(row)] for row in _priority_rows(run)]))
    parts.append('<h2 id="gaps">待補資料與證據清單</h2><p>以下項目需要補充規格、適用條件、原文或測試報告；缺少證據不直接等同產品不符合。</p>')
    parts.append(_table(["結果／規範", "要求", "目前判定", "缺漏與待釐清", "建議補件／確認動作", "既有原文證據"], [[f"{row.get('id', '')}\n{row.get('standard_name', '')}", row.get("requirement", ""), STATUSES.get(_effective(row), _effective(row)), "\n".join([row.get("explanation") or ""] + (row.get("differences") or []) + (row.get("warnings") or [])), _effective_risk(row).get("action", ""), _evidence_text(row.get("evidence") or [])] for row in _evidence_gaps(run)]))
    parts.append('<h2 id="results">全部項目差異與人工覆核</h2>')
    for index, row in enumerate(run.get("results", []), 1):
        review = row.get("review") or {}
        parts.append(f'<article class="row"><h3>{index}. {_escape(row.get("standard_name", ""))} — {_escape((row.get("item") or {}).get("name") or row.get("location", ""))}</h3>')
        parts.append(_paragraph(f"結果 ID：{row.get('id', '')}　項目 ID：{_item_id(row)}　位置：{row.get('location', '')}", "muted"))
        parts.append(_paragraph(f"目前判定：{STATUSES.get(_effective(row), _effective(row))}　AI 判定：{STATUSES.get(row.get('status'), row.get('status', ''))}\n目前風險：{_risk_label(_effective_risk(row).get('level'))}　AI 風險：{_risk_label(_risk_value(row.get('ai_risk') or row.get('risk')).get('level'))}"))
        parts.append("<h4>規範項目與原文</h4>" + _paragraph(_item_summary(row.get("item") or {})) + _paragraph(row.get("requirement", "")))
        parts.append("<h4>判定說明與全部差異</h4>" + _paragraph(row.get("explanation", "")) + _list(row.get("differences") or []))
        parts.append("<h4>風險依據與建議動作</h4>" + _paragraph(_risk_summary(row)) + _paragraph(_json({"effective": _effective_risk(row), "ai_risk": row.get("ai_risk", {})}), "metadata"))
        parts.append("<h4>產品原文證據</h4>")
        for evidence in row.get("evidence") or []:
            parts.append('<div class="evidence">' + _paragraph(f"{evidence.get('block_id', '')}　{evidence.get('location', '')}", "muted") + _paragraph(evidence.get("quote", "")) + "</div>")
        if not row.get("evidence"):
            parts.append("<p>無有效原文引述。</p>")
        parts.append("<h4>檢索覆蓋、對應產品項目與警告</h4>" + _paragraph(_json({"retrieval": row.get("retrieval", {}), "matched_product_item_ids": row.get("matched_product_item_ids", []), "product_coverage": row.get("product_coverage", {})}), "metadata") + _list(row.get("warnings") or []))
        parts.append("<h4>目前人工覆核</h4>" + _paragraph(f"狀態：{DECISIONS.get(review.get('decision', 'pending'), review.get('decision', 'pending'))}\n覆核人員：{review.get('reviewer', '')}　版本：{review.get('version', 0)}　時間：{review.get('updated_at') or ''}\n人工判定：{STATUSES.get(review.get('final_status'), review.get('final_status') or '')}　人工風險：{_risk_label(review.get('risk_level'))}\n註記：{review.get('note', '')}"))
        parts.append("<h4>全部覆核事件</h4>" + _paragraph(_json(row.get("history", [])), "metadata") + "</article>")
    extras, extras_metadata = _extra_items(run)
    parts.append('<h2 id="extras">尚未建立對應的產品項目</h2><p class="notice">未對應不代表標準沒有要求。結果受選取範圍、抽取與檢索覆蓋限制；只有已處理項目中的對應關係已被記錄。</p>' + _paragraph(_json({key: value for key, value in extras_metadata.items() if key != "items"}), "metadata"))
    parts.append(_table(["項目 ID", "產品項目", "來源原文", "條件與完整資料"], [[item.get("id"), item.get("name"), item.get("quote"), _item_summary(item)] for item in extras]))
    parts.append('<h2 id="sources">不可變文件快照、抽取項目與來源</h2>')
    for document in run.get("documents", []):
        metadata = {key: value for key, value in document.items() if key not in {"blocks", "items", "extraction_audit"}}
        parts.append(f'<section class="source"><h3>{_escape(document.get("name", ""))}</h3>' + _paragraph(_json(metadata), "metadata"))
        for block in document.get("blocks", []):
            parts.append(f'<h4>{_escape(block.get("id", ""))} — {_escape(block.get("location", ""))}</h4>' + _paragraph(block.get("text", "")))
        parts.append("<h4>全部抽取項目</h4>" + _paragraph(_json(document.get("items", [])), "metadata"))
        parts.append("<h4>抽取與項目修訂歷程</h4>" + _paragraph(_json(document.get("extraction_audit", [])), "metadata") + "</section>")
    parts.append('<h2 id="audit">分析、選取與覆核完整操作歷程</h2>' + _paragraph(_json(run.get("audit", [])), "metadata"))
    parts.append("<h2>分析設定與來源中繼資料（不含密鑰）</h2>" + _paragraph(_json({key: value for key, value in run.items() if key not in {"documents", "results", "screenings", "product_extras", "audit", "rankings"}}), "metadata"))
    parts.append("<h2>完整分析資料快照（JSON）</h2>" + _paragraph(_json(run), "metadata") + "</body></html>")
    return "".join(parts).encode("utf-8")


def _html(run: dict) -> bytes:
    sections = ["""<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>規格比對與人工覆核報告</title>
<style>body{font-family:system-ui,'Microsoft JhengHei',sans-serif;color:#172b3a;max-width:1200px;margin:32px auto;padding:0 24px;line-height:1.65}h1,h2,h3{line-height:1.3}h2{border-bottom:2px solid #17696a;padding-bottom:8px;margin-top:40px}h3{margin-top:12px}pre{font-family:inherit;white-space:pre-wrap;overflow-wrap:anywhere;margin:8px 0}table{border-collapse:collapse;width:100%;margin:12px 0}td,th{border:1px solid #c8d5dc;padding:8px;text-align:left;vertical-align:top;overflow-wrap:anywhere}th{background:#edf4f6}.row,.source{border:1px solid #c8d5dc;border-radius:8px;padding:18px;margin:16px 0}.muted{color:#526571;font-size:.9em}.notice{background:#fff4d5;border-left:4px solid #ad7700;padding:14px}.badge{display:inline-block;background:#e8f3f2;padding:3px 12px;border-radius:12px}.evidence{border-left:3px solid #408787;padding-left:14px}.metadata{font-size:.85em}a{color:#136e70}@media print{body{max-width:none;margin:0;padding:0;font-size:10pt}.row,.source{border-radius:0}h2,h3{break-after:avoid}thead{display:table-header-group}pre,td{overflow-wrap:anywhere}.row{break-inside:auto}@page{margin:16mm}}</style></head><body>"""]
    sections.append("<h1>規格比對與人工覆核報告</h1>")
    sections.append(_paragraph(f"執行 ID：{run.get('id', '')}　狀態：{run.get('status', '')}\n建立時間：{run.get('created_at', '')}　完成時間：{run.get('finished_at') or '尚未完成'}\n模式：{run.get('mode', '')}　完成區塊：{run.get('completed', 0)} / {run.get('total', 0)}"))
    notice = "此報告保留已產生的全部比對列、原文證據、文件抽取快照及覆核歷程。文件符合度參考分數不代表法規認證；區塊覆蓋不保證所有語意差異已被找出。產品文件未載明不等於產品本身不合規。覆核人員為自行填寫，非驗證身分或電子簽章。"
    if run.get("status") != "completed":
        notice += " 此次執行尚未完整完成，分數及結果屬暫定；未產生結果的區塊仍需比對。"
    if run.get("mode") == "demo":
        notice += " 這是合成範例示範結果，未呼叫本地模型，不可作為真實產品判定。"
    sections.append(f'<p class="notice">{_escape(notice)}</p>')
    if run.get("error"):
        sections.append("<h3>執行診斷</h3>" + _paragraph(run["error"]))
    sections.append('<p><a href="#results">逐條差異與覆核</a> · <a href="#sources">原始文件文字快照</a> · <a href="#run">執行設定</a></p>')
    sections.append("<h2>文件符合度參考分數</h2><table><thead><tr><th>規範文件</th><th>參考分數</th><th>完成／總區塊</th><th>已覆核</th><th>各狀態與覆蓋率</th></tr></thead><tbody>")
    for ranking in run.get("rankings", []):
        score = ranking.get("score", 0)
        score_text = f"{score:.2f}" if isinstance(score, (int, float)) else str(score)
        sections.append("<tr>" + "".join(f"<td>{_escape(value)}</td>" for value in [ranking.get("standard_name", ""), score_text, f"{ranking.get('completed', 0)} / {ranking.get('total', 0)}", ranking.get("reviewed", 0), _json({"counts": ranking.get("counts", {}), "coverage": ranking.get("coverage")})]) + "</tr>")
    sections.append("</tbody></table><p class=muted>分數＝（符合＋0.5 × 部分符合）÷ 全部規範區塊 × 100；採用已確認／更正的人工判定，其餘採 AI 判定。</p>")
    sections.append('<h2 id="results">逐條差異與人工覆核</h2>')
    for index, row in enumerate(run.get("results", []), 1):
        review = row.get("review") or {}
        sections.append(f'<article class="row"><h3>{index}. {_escape(row.get("standard_name", ""))} — {_escape(row.get("location", ""))}</h3>')
        sections.append(f'<p><span class="badge">目前判定：{_escape(STATUSES.get(_effective(row), _effective(row)))}</span>　AI：{_escape(STATUSES.get(row.get("status"), row.get("status", "")))}　信心值：{_escape(row.get("confidence", ""))}</p>')
        sections.append(_paragraph(f"結果 ID：{row.get('id', '')}　規範 ID：{row.get('standard_id', '')}　區塊：{row.get('block_id', '')}", "muted"))
        sections.append("<h4>規範原文</h4>" + _paragraph(row.get("requirement", "")))
        sections.append("<h4>判定說明</h4>" + _paragraph(row.get("explanation", "")))
        sections.append("<h4>差異點</h4>" + _list(row.get("differences") or []))
        sections.append("<h4>產品原文證據</h4>")
        for evidence in row.get("evidence") or []:
            sections.append('<div class="evidence">' + _paragraph(f"{evidence.get('block_id', '')}　{evidence.get('location', '')}", "muted") + _paragraph(evidence.get("quote", "")) + "</div>")
        if not row.get("evidence"):
            sections.append("<p>無有效原文引述。</p>")
        sections.append("<h4>抽取／比對警告與產品覆蓋</h4>" + _list(row.get("warnings") or []) + _paragraph(_json(row.get("product_coverage", {})), "metadata"))
        sections.append("<h4>目前覆核</h4>" + _paragraph(f"狀態：{DECISIONS.get(review.get('decision', 'pending'), review.get('decision', 'pending'))}\n覆核人員：{review.get('reviewer', '')}\n人工判定：{STATUSES.get(review.get('final_status'), review.get('final_status') or '')}\n更新：{review.get('updated_at') or ''}　版本：{review.get('version', 0)}\n註記：{review.get('note', '')}"))
        sections.append("<h4>完整覆核歷程</h4>" + _paragraph(_json(row.get("history", [])), "metadata") + "</article>")
    sections.append('<h2 id="sources">原始文件文字快照與來源資訊</h2><p>以下保留本次執行使用的全部抽取區塊。圖像與抽取限制請參考各文件警告；原始上傳檔另由系統保存。</p>')
    for document in run.get("documents", []):
        metadata = {key: value for key, value in document.items() if key != "blocks"}
        sections.append(f'<section class="source"><h3>{_escape(document.get("name", ""))}</h3>' + _paragraph(_json(metadata), "metadata"))
        for block in document.get("blocks", []):
            sections.append(f'<h4>{_escape(block.get("id", ""))} — {_escape(block.get("location", ""))}</h4>' + _paragraph(block.get("text", "")))
        sections.append("</section>")
    sections.append('<h2 id="run">執行設定與中繼資料（不含密鑰）</h2>')
    sections.append(_paragraph(_json({key: value for key, value in run.items() if key not in {"documents", "results", "rankings"}}), "metadata"))
    # Keep unknown extension fields too, without exposing executable script content.
    sections.append("<h2>完整執行資料快照（JSON）</h2>" + _paragraph(_json(run), "metadata"))
    sections.append("</body></html>")
    return "".join(sections).encode("utf-8")


def _excel_text(value) -> str:
    text = _json(value) if isinstance(value, (dict, list, tuple)) else str(value if value is not None else "")
    # XML 1.0 disallows these; retain an explicit visible encoding instead of dropping.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", lambda match: f"\\u{ord(match.group(0)):04x}", text)


def _xlsx(run: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name: str, headings: list[str]):
        current = workbook.create_sheet(name)
        current.append(headings + ["分段"])
        current.freeze_panes = "A2"
        for cell in current[1]:
            cell.fill = PatternFill("solid", fgColor="17696A")
            cell.font = Font(color="FFFFFF", bold=True)
            current.column_dimensions[cell.column_letter].width = 24
        return current

    def append(current, values: list):
        # Excel silently truncates cells >32767 chars; use explicit continuation rows.
        texts = [_excel_text(value) for value in values]
        pieces = [[text[index:index + 30_000] for index in range(0, len(text), 30_000)] or [""] for text in texts]
        count = max(len(part) for part in pieces)
        for index in range(count):
            row = [part[index] if index < len(part) else "" for part in pieces]
            row.append(f"{index + 1}/{count}")
            current.append(row)
            for cell in current[current.max_row]:
                # Explicit text cells prevent =,+,-,@ from becoming executable formulas.
                cell.data_type = "s"
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    analysis = _is_analysis(run)
    results = sheet("Results", ["結果 ID", "規範 ID", "規範文件", "區塊", "位置", "規範原文", "AI 判定", "目前判定", "信心值", "說明", "全部差異點 JSON", "全部產品證據 JSON", "產品覆蓋 JSON", "警告 JSON", "覆核狀態", "人工判定", "覆核人員", "覆核註記", "覆核時間", "覆核版本", "完整結果 JSON"] + (["目前風險", "AI 風險", "項目 ID"] if analysis else []))
    history = sheet("History", ["結果 ID", "事件 ID", "決定", "最終狀態", "覆核人員", "註記", "時間", "版本", "完整事件 JSON"])
    for row in run.get("results", []):
        review = row.get("review") or {}
        append(results, [row.get("id"), row.get("standard_id"), row.get("standard_name"), row.get("block_id"), row.get("location"), row.get("requirement"), row.get("status"), _effective(row), row.get("confidence"), row.get("explanation"), row.get("differences", []), row.get("evidence", []), row.get("product_coverage", {}), row.get("warnings", []), review.get("decision"), review.get("final_status"), review.get("reviewer"), review.get("note"), review.get("updated_at"), review.get("version", 0), row] + ([_effective_risk(row).get("level"), _risk_value(row.get("ai_risk") or row.get("risk")).get("level"), _item_id(row)] if analysis else []))
        for event in row.get("history", []):
            append(history, [row.get("id"), event.get("id"), event.get("decision"), event.get("final_status"), event.get("reviewer"), event.get("note"), event.get("created_at"), event.get("version"), event])
    documents = sheet("Documents", ["文件 ID", "檔名", "角色", "SHA256", "建立時間", "抽取已確認", "警告 JSON", "區塊 ID", "來源位置", "全部抽取文字", "文件中繼資料 JSON"])
    for document in run.get("documents", []):
        # Analysis items/audits have dedicated sheets: repeating them once per
        # source block can multiply a large library into gigabytes of cells.
        separate = {"blocks", "items", "extraction_audit"} if analysis else {"blocks"}
        metadata = {key: value for key, value in document.items() if key not in separate}
        for block in document.get("blocks") or [{}]:
            append(documents, [document.get("id"), document.get("name"), document.get("role"), document.get("sha256"), document.get("created_at"), document.get("confirmed", document.get("extraction_confirmed")), document.get("warnings", []), block.get("id"), block.get("location"), block.get("text"), metadata])
    metadata_sheet = sheet("Run", ["欄位", "值"])
    append(metadata_sheet, ["exported_at", datetime.now(timezone.utc).isoformat()])
    append(metadata_sheet, ["說明", "相關性、適用性、符合判定與風險分別呈現。風險為覆核優先提示，不是經校準的事故機率或法規認證。結構覆蓋不保證語意完整性。產品文件未載明不等於產品不合規。未對應產品不代表標準沒有要求。覆核身分為自行填寫，非電子簽章。" if analysis else "文件符合度參考分數不是法規認證。區塊覆蓋不保證全部語意差異已找出。產品文件未載明不等於產品不合規。未完整執行的結果與分數屬暫定。覆核身分為自行填寫，非電子簽章。示範模式僅供展示。"])
    append(metadata_sheet, ["長文字與安全", "所有儲存格均為純文字；超過 30,000 字元拆為連續分段列，續列順序依「分段」欄。XML 不允許的控制字元以 \\uXXXX 顯示。完整結果 JSON 可依分段順序還原。"])
    for key, value in run.items():
        if key not in {"documents", "results"}:
            append(metadata_sheet, [key, value])
    if analysis:
        append(metadata_sheet, ["partial_analysis", _analysis_partial(run)])
        append(metadata_sheet, ["覆蓋說明", "結構覆蓋是已處理文件、區塊、項目或檢索分段數，不保證全部語意差異已找到。聚焦比對未檢索部分及尚未處理項目仍需確認。"])
        screenings = sheet("Screenings", ["初篩 ID", "標準 ID", "標準名稱", "相關性", "適用性", "理由", "原文證據 JSON", "警告 JSON", "本次選取", "人工選取覆核 JSON", "完整初篩 JSON"])
        selection_audit = sheet("SelectionAudit", ["來源", "標準 ID", "標準名稱", "本次選取", "完整選取事件 JSON"])
        for entry in run.get("screenings", []):
            append(screenings, [entry.get("id"), entry.get("standard_id"), entry.get("standard_name"), entry.get("relevance"), entry.get("applicability"), entry.get("reason"), entry.get("evidence", []), entry.get("warnings", []), entry.get("selected"), entry.get("review", {}), entry])
            append(selection_audit, ["screening.review", entry.get("standard_id"), entry.get("standard_name"), entry.get("selected"), entry.get("review", {})])
            for event in entry.get("history", []):
                append(selection_audit, ["screening.history", entry.get("standard_id"), entry.get("standard_name"), event.get("selected"), event])
        risks = sheet("Risks", ["結果 ID", "標準 ID", "標準名稱", "項目 ID", "項目名稱", "目前判定", "目前風險", "AI 風險", "提示類型", "依據", "建議動作", "人工風險", "覆核人員", "覆核原因", "完整風險 JSON", "完整 AI 風險 JSON"])
        for row in _priority_rows(run):
            risk = _effective_risk(row)
            review = row.get("review") or {}
            append(risks, [row.get("id"), row.get("standard_id"), row.get("standard_name"), _item_id(row), (row.get("item") or {}).get("name"), _effective(row), risk.get("level"), _risk_value(row.get("ai_risk") or row.get("risk")).get("level"), risk.get("kind"), risk.get("basis"), risk.get("action"), review.get("risk_level"), review.get("reviewer"), review.get("note"), risk, row.get("ai_risk", {})])
        gaps = sheet("EvidenceGaps", ["結果 ID", "標準名稱", "項目 ID", "要求原文", "目前判定", "說明", "差異 JSON", "建議動作", "產品證據 JSON", "檢索覆蓋 JSON", "警告 JSON"])
        for row in _evidence_gaps(run):
            append(gaps, [row.get("id"), row.get("standard_name"), _item_id(row), row.get("requirement"), _effective(row), row.get("explanation"), row.get("differences", []), _effective_risk(row).get("action"), row.get("evidence", []), row.get("retrieval", {}), row.get("warnings", [])])
        extras = sheet("ProductExtras", ["項目 ID", "產品項目", "原文", "來源區塊", "來源位置", "完整覆蓋", "限制說明", "完整項目 JSON"])
        extra_items, extra_metadata = _extra_items(run)
        extra_notice = "尚未對應不代表標準沒有要求；受本次選取、抽取及檢索覆蓋限制。"
        for item in extra_items:
            append(extras, [item.get("id"), item.get("name"), item.get("quote"), item.get("block_id"), item.get("location"), extra_metadata.get("coverage_complete", False), extra_notice + str(extra_metadata.get("note") or ""), item])
        # Include the limitation even when no unmatched item has been produced.
        append(metadata_sheet, ["product_extras_notice", extra_notice])
        items = sheet("Items", ["文件 ID", "文件名稱", "角色", "項目 ID", "項目名稱", "種類", "參數", "數值", "單位", "運算子", "條件", "例外", "測試方法", "重要性", "重要性依據", "原文", "區塊 ID", "來源位置", "完整項目 JSON"])
        extraction_audit = sheet("ExtractionAudit", ["文件 ID", "文件名稱", "事件順序", "完整抽取／修訂事件 JSON"])
        for document in run.get("documents", []):
            for item in document.get("items", []):
                append(items, [document.get("id"), document.get("name"), document.get("role"), item.get("id"), item.get("name"), item.get("kind"), item.get("parameter"), item.get("value"), item.get("unit"), item.get("operator"), item.get("conditions"), item.get("exceptions"), item.get("test_method"), item.get("criticality"), item.get("criticality_basis"), item.get("quote"), item.get("block_id"), item.get("location"), item])
            for index, event in enumerate(document.get("extraction_audit", []), 1):
                append(extraction_audit, [document.get("id"), document.get("name"), index, event])
        audit = sheet("Audit", ["事件順序", "事件 ID", "事件類型", "時間", "完整操作事件 JSON"])
        for index, event in enumerate(run.get("audit", []), 1):
            event_type = event.get("type") or event.get("action") or event.get("kind") or ""
            append(audit, [index, event.get("id"), event_type, event.get("created_at") or event.get("at"), event])
            if any(token in str(event_type).lower() for token in ("select", "exclud", "選取", "排除")):
                append(selection_audit, ["analysis.audit", event.get("standard_id"), event.get("standard_name"), event.get("selected"), event])
        unprocessed = sheet("Unprocessed", ["標準 ID", "標準名稱", "項目 ID", "項目名稱", "原文", "狀態", "完整項目 JSON"])
        for entry in _unprocessed_items(run):
            item = entry["item"]
            append(unprocessed, [entry.get("standard_id"), entry.get("standard_name"), item.get("id"), item.get("name"), item.get("quote"), "未處理／無結果", item])
    for current in workbook.worksheets:
        current.auto_filter.ref = current.dimensions
    data = io.BytesIO()
    workbook.save(data)
    return data.getvalue()


def export_run(run: dict, format: str) -> tuple[bytes, str, str]:
    """Return report bytes, content type and a safe attachment filename."""
    if format not in {"html", "xlsx", "json"}:
        raise ValueError("匯出格式必須是 html、xlsx 或 json。")
    snapshot = _sanitize(run)
    identifier = re.sub(r"[^A-Za-z0-9_-]", "_", str(run.get("id", "report")))[:80] or "report"
    name = f"spec-check-{identifier}.{format}"
    if format == "json":
        return _json(snapshot).encode("utf-8"), "application/json; charset=utf-8", name
    if format == "html":
        return (_analysis_html(snapshot) if _is_analysis(snapshot) else _html(snapshot)), "text/html; charset=utf-8", name
    return _xlsx(snapshot), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", name
