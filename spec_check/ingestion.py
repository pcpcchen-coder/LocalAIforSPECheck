"""Local, traceable text extraction; never silently omit oversized text."""

from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree

MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_ARCHIVE_BYTES = 128 * 1024 * 1024
MAX_MEMBER_BYTES = 32 * 1024 * 1024
MAX_BLOCK_CHARS = 1400
MAX_EXTRACTED_CHARS = 12_000_000
MAX_BLOCKS = 30_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md"}
csv.field_size_limit(MAX_FILE_BYTES)


class _Collector:
    def __init__(self) -> None:
        self.blocks: list[dict] = []
        self.warnings: list[str] = []
        self.character_count = 0

    def warn(self, message: str) -> None:
        if message not in self.warnings:
            self.warnings.append(message)

    def add(self, text: str, location: str) -> None:
        if not text or not text.strip():
            return
        # Preserve source characters, including indentation and internal whitespace.
        self.character_count += len(text)
        if self.character_count > MAX_EXTRACTED_CHARS:
            raise ValueError("抽取文字超過 1,200 萬字元，請拆分文件後重新上傳；本次未截斷保存。")
        pieces = []
        remaining = text
        while remaining:
            end = min(MAX_BLOCK_CHARS, len(remaining))
            if end < len(remaining):
                newline = remaining.rfind("\n", 0, end)
                if newline >= 0:
                    end = newline + 1
            pieces.append(remaining[:end])
            remaining = remaining[end:]
        for index, piece in enumerate(pieces, 1):
            if not piece.strip():
                continue
            if len(self.blocks) >= MAX_BLOCKS:
                raise ValueError("抽取區塊超過 30,000 個，請拆分文件後重新上傳；本次未截斷保存。")
            suffix = f"（分段 {index}/{len(pieces)}）" if len(pieces) > 1 else ""
            self.blocks.append({"id": f"B{len(self.blocks) + 1:05d}", "location": location + suffix, "text": piece})


def _check_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10_000 or sum(item.file_size for item in entries) > MAX_ARCHIVE_BYTES:
                raise ValueError("壓縮文件內容過大（最多 10,000 項、解壓總量 128 MB）。")
            for item in entries:
                if item.file_size > MAX_MEMBER_BYTES:
                    raise ValueError("壓縮文件單一項目超過 32 MB，請縮小文件。")
                if item.file_size > max(item.compress_size, 1) * 1000:
                    raise ValueError("壓縮文件的壓縮比例異常，已停止處理。")
                if item.flag_bits & 1:
                    raise ValueError("不支援加密的 Office 文件，請先另存未加密副本。")
    except zipfile.BadZipFile as exc:
        raise ValueError("Office 文件不是有效的 DOCX/XLSX 壓縮格式。") from exc


def _read_pdf(path: Path, output: _Collector) -> None:
    from pypdf import PdfReader

    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF 需要密碼，請先另存未加密副本。")
    if len(reader.pages) > 3000:
        raise ValueError("PDF 超過 3,000 頁，請拆分文件後重新上傳。")
    output.warn("PDF 僅抽取文字層，未執行 OCR；表格欄位、圖形、公式、雙欄閱讀順序及字型編碼可能失真，請逐頁對照原檔確認。")
    for index, page in enumerate(reader.pages, 1):
        if _pdf_has_images(page.get("/Resources")):
            output.warn(f"PDF 第 {index} 頁包含圖像；圖中文字／規格未 OCR，即使同頁有文字層也可能漏掉圖像內容。")
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            output.warn(f"PDF 第 {index} 頁文字抽取失敗（{type(exc).__name__}），未涵蓋該頁；請人工確認或提供文字版。")
            continue
        if not text.strip():
            output.warn(f"PDF 第 {index} 頁沒有可抽取文字，可能為掃描／圖片或空白頁；本工具未 OCR，該頁不可視為已完成比對。")
        pending = ""
        first_line = 1
        last_line = 0
        for line_number, line in enumerate(text.splitlines(keepends=True), 1):
            if pending and len(pending) + len(line) > MAX_BLOCK_CHARS:
                output.add(pending, f"PDF 第 {index} 頁 第 {first_line}–{last_line} 行")
                pending = ""
            if not pending:
                first_line = line_number
            pending += line
            last_line = line_number
        if pending:
            output.add(pending, f"PDF 第 {index} 頁 第 {first_line}–{last_line} 行")


def _pdf_has_images(resources, seen=None) -> bool:
    """Inspect image resources without decoding or executing image payloads."""
    if not resources:
        return False
    seen = set() if seen is None else seen
    try:
        resources = resources.get_object()
        marker = id(resources)
        if marker in seen:
            return False
        seen.add(marker)
        if len(seen) > 1000:
            return True
        objects = resources.get("/XObject")
        if not objects:
            return False
        for reference in objects.get_object().values():
            item = reference.get_object()
            if item.get("/Subtype") == "/Image" or _pdf_has_images(item.get("/Resources"), seen):
                return True
    except Exception:
        # Unknown resources remain covered by the generic PDF extraction warning.
        return False
    return False


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _xml_paragraph_text(element) -> str:
    pieces = []
    for node in element.iter():
        if node.tag == _W + "t":
            pieces.append(node.text or "")
        elif node.tag == _W + "tab":
            pieces.append("\t")
        elif node.tag in {_W + "br", _W + "cr"}:
            pieces.append("\n")
    return "".join(pieces)


def _read_docx(path: Path, output: _Collector) -> None:
    from docx import Document

    document = Document(str(path))
    output.warn("DOCX 已抽取本文、表格及頁首／頁尾等文字；圖片、繪圖、公式、自動編號（條號）、修訂顯示狀態與視覺版面未可靠解析，請對照原檔確認。")
    output.warn("DOCX 表格以第一個非空列作為推定表頭，附於後續列以保留欄位上下文；多層表頭、巢狀表格、跨列／跨欄合併需人工確認。")
    paragraph_index = table_index = 0
    for child in document.element.body:
        if child.tag == _W + "p":
            paragraph_index += 1
            output.add(_xml_paragraph_text(child), f"DOCX 本文段落 {paragraph_index}")
        elif child.tag == _W + "tbl":
            table_index += 1
            header = ""
            # Each row remains one logical source; nested paragraphs/tables are retained.
            for row_index, row in enumerate(child.findall(_W + "tr"), 1):
                cells = []
                for cell_index, cell in enumerate(row.findall(_W + "tc"), 1):
                    paragraphs = [_xml_paragraph_text(p) for p in cell.iter(_W + "p")]
                    cells.append(f"欄 {cell_index}：" + "\n".join(paragraphs))
                if any(cell.split("：", 1)[1].strip() for cell in cells):
                    row_text = " | ".join(cells)
                    text = f"推定表頭原文：{header}\n本列原文：{row_text}" if header else row_text
                    output.add(text, f"DOCX 表格 {table_index} 第 {row_index} 列")
                    if not header:
                        header = row_text
        elif child.tag != _W + "sectPr":
            # Content controls and other wrappers may contain meaningful paragraphs.
            for paragraph in child.iter(_W + "p"):
                paragraph_index += 1
                output.add(_xml_paragraph_text(paragraph), f"DOCX 包裝區段段落 {paragraph_index}")
    labels = {"header": "頁首", "footer": "頁尾", "footnotes": "註腳", "endnotes": "章節附註", "comments": "批註"}
    with zipfile.ZipFile(path) as archive:
        for filename in sorted(archive.namelist()):
            match = re.fullmatch(r"word/(header\d+|footer\d+|footnotes|endnotes|comments)\.xml", filename)
            if not match:
                continue
            kind = re.sub(r"\d+$", "", match.group(1))
            root = ElementTree.fromstring(archive.read(filename))
            for index, paragraph in enumerate(root.iter(_W + "p"), 1):
                output.add(_xml_paragraph_text(paragraph), f"DOCX {labels[kind]} {match.group(1)} 段落 {index}")
            if kind in {"footnotes", "endnotes", "comments"}:
                output.warn(f"已納入 DOCX {labels[kind]}文字，請判斷其是否屬於正式規格要求。")


def _cell_text(value) -> str:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _check_sheet_bounds(sheet) -> tuple[int, int, int, int]:
    """Check actual XML coordinates before openpyxl can allocate sparse rows.

    The worksheet's optional dimension declaration can be stale or incorrect and
    must never determine extraction coverage or resource limits.
    """
    from openpyxl.utils.cell import column_index_from_string

    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    current_row = max_row = max_column = previous_column = 0
    min_row = min_column = None
    seen_columns: set[int] = set()
    with sheet._get_source() as source:
        for event, element in ElementTree.iterparse(source, events=("start", "end")):
            if event == "end":
                element.clear()
                continue
            if element.tag == namespace + "row":
                row_text = element.get("r", str(current_row + 1))
                if not row_text.isdigit() or len(row_text) > 7:
                    raise ValueError(f"工作表「{sheet.title}」包含無效列座標，請重新另存文件。")
                row = int(row_text)
                if row <= current_row:
                    raise ValueError(f"工作表「{sheet.title}」列座標重複或順序錯誤，請重新另存文件。")
                current_row = row
                min_row = row if min_row is None else min(min_row, row)
                max_row = max(max_row, row)
                previous_column = 0
                seen_columns.clear()
            elif element.tag == namespace + "c":
                coordinate = element.get("r")
                column = previous_column + 1
                if coordinate:
                    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", coordinate)
                    if not match or int(match.group(2)) != current_row:
                        raise ValueError(f"工作表「{sheet.title}」包含無效儲存格座標，請重新另存文件。")
                    column = column_index_from_string(match.group(1))
                if not current_row or column in seen_columns or column > 16_384:
                    raise ValueError(f"工作表「{sheet.title}」儲存格座標重複或無效，請重新另存文件。")
                seen_columns.add(column)
                previous_column = column
                max_column = max(max_column, column)
                min_column = column if min_column is None else min(min_column, column)
            if max_row * max(max_column, 1) > 1_000_000:
                raise ValueError(f"工作表「{sheet.title}」實際使用範圍超過 100 萬格，請縮小範圍或拆分文件。")
    return min_row or 1, min_column or 1, max_row, max_column


def _read_xlsx(path: Path, output: _Collector) -> None:
    from openpyxl import load_workbook

    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    output.warn("XLSX 依工作表／列抽取儲存格，包含隱藏工作表；圖片、圖表、樣式、批註及合併儲存格的視覺關係需人工確認。本工具不重算公式。")
    output.warn("XLSX 以每張工作表第一個非空列作為推定表頭，附於後續列以保留欄位上下文；多層表頭及跨列／跨欄合併需人工確認。")
    try:
        for sheet in formulas.worksheets:
            actual_bounds = _check_sheet_bounds(sheet)
            if actual_bounds[2] and (sheet.min_row, sheet.min_column, sheet.max_row, sheet.max_column) != actual_bounds:
                output.warn(f"工作表「{sheet.title}」宣告範圍與實際列／儲存格不一致，已忽略宣告並依實際內容完整抽取。")
            # Both parsers must ignore potentially undersized dimension metadata,
            # otherwise either formulas or cached values could silently disappear.
            sheet.reset_dimensions()
            cached_sheet = cached[sheet.title]
            cached_sheet.reset_dimensions()
            if sheet.sheet_state != "visible":
                output.warn(f"已納入隱藏工作表「{sheet.title}」，請確認是否屬於正式規格。")
            cached_rows = cached_sheet.iter_rows(min_row=1, min_col=1)
            header = ""
            for row_index, row in enumerate(sheet.iter_rows(min_row=1, min_col=1), 1):
                value_row = next(cached_rows, ())
                cells = []
                for cell_index, cell in enumerate(row):
                    if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
                        continue
                    value = _cell_text(cell.value)
                    if cell.data_type == "f":
                        cached_value = value_row[cell_index].value if cell_index < len(value_row) else None
                        if cached_value is None:
                            output.warn(f"工作表「{sheet.title}」{cell.coordinate} 公式沒有可讀取的快取結果，請在 Excel／LibreOffice 重算並另存後上傳；目前僅保留公式。")
                            value += " [公式結果未快取]"
                        else:
                            value = f"{_cell_text(cached_value)} [公式：{value}]"
                    if cell.number_format and cell.number_format != "General":
                        value += f" [顯示格式：{cell.number_format}]"
                    cells.append(f"{cell.coordinate}：{value}")
                if cells:
                    row_text = " | ".join(cells)
                    text = f"推定表頭原文：{header}\n本列原文：{row_text}" if header else row_text
                    output.add(text, f"XLSX 工作表「{sheet.title}」第 {row_index} 列")
                    if not header:
                        header = row_text
    finally:
        formulas.close()
        cached.close()


def _read_text(path: Path, extension: str, output: _Collector) -> None:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("文字／CSV 文件必須為 UTF-8（可含 BOM），請轉存 UTF-8 後上傳。") from exc
    if "\x00" in text:
        raise ValueError("文字檔含有 NUL 字元，可能不是 UTF-8 純文字文件。")
    if extension == ".csv":
        # utf-8 and CSV's own quoting rules preserve multiline cell content.
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        previous_line = 0
        header = ""
        output.warn("CSV 第一個非空列會作為推定表頭，附於後續列；若文件沒有表頭，請在抽取預覽確認上下文。")
        try:
            for row_index, row in enumerate(reader, 1):
                first_line = previous_line + 1
                previous_line = reader.line_num
                if any(value.strip() for value in row):
                    row_text = " | ".join(f"欄 {index}：{value}" for index, value in enumerate(row, 1))
                    rendered = f"推定表頭原文：{header}\n本列原文：{row_text}" if header else row_text
                    output.add(rendered, f"CSV 第 {row_index} 列（原檔第 {first_line}–{reader.line_num} 行）")
                    if not header:
                        header = row_text
        except csv.Error as exc:
            raise ValueError(f"CSV 格式無法解析：{exc}") from exc
    else:
        for line_number, line in enumerate(text.splitlines(), 1):
            output.add(line, f"文字第 {line_number} 行")


def extract_document(path: Path, name: str, role: str) -> dict:
    """Return all extractable text with source locations, limitations and hash."""
    path = Path(path)
    if role not in {"product", "standard"}:
        raise ValueError("文件角色必須是 product 或 standard。")
    extension = Path(name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("不支援此格式。請使用 PDF、DOCX、XLSX、CSV、TXT 或 MD。")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("單一文件不得超過 30 MB，請拆分後上傳。")
    if extension in {".docx", ".xlsx"}:
        _check_archive(path)
    output = _Collector()
    output.warn("比對區塊依文件結構與 1,400 字元上限分段，不保證等同完整、獨立的規範條款；跨段條件、章節適用範圍及表頭仍須人工覆核。")
    try:
        if extension == ".pdf":
            _read_pdf(path, output)
        elif extension == ".docx":
            _read_docx(path, output)
        elif extension == ".xlsx":
            _read_xlsx(path, output)
        else:
            _read_text(path, extension, output)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"無法解析 {extension.upper()} 文件（{type(exc).__name__}）。請確認檔案完整且未加密。") from exc
    if not output.blocks:
        output.warn("沒有抽取到任何文字，無法執行規格比對。請提供文字版文件或先完成 OCR。")
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return {"name": name, "role": role, "sha256": digest.hexdigest(), "blocks": output.blocks, "warnings": output.warnings}
